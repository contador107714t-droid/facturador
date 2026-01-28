from fastapi import FastAPI, Request
from fastapi import Form
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware
import sqlite3
import os
from passlib.context import CryptContext
from db_auth import get_user_by_username
import datetime
from urllib.parse import quote
import json
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

app = FastAPI()
templates = Jinja2Templates(directory="templates")
def inject_user(request: Request):
    return {"user": request.session.get("user")}

templates.env.globals["current_user"] = inject_user

app.mount("/static", StaticFiles(directory="static"), name="static")
pwd_context = CryptContext(schemes=["pbkdf2_sha256","bcrypt"], deprecated="auto")


PUBLIC_PATHS = ("/login", "/static", "/favicon.ico")

@app.middleware("http")
async def require_login(request: Request, call_next):
    path = request.url.path

    # Permite login, logout y archivos estáticos
    if path in ("/login", "/logout") or path.startswith("/static"):
        return await call_next(request)


    # Restricción de permisos para role cliente_receptor
    user = request.session.get("user") or {}
    if user.get("role") == "cliente_receptor":
        if path == "/emisores" or path.startswith("/emisores/") or path == "/receptores" or path.startswith("/receptores/"):
            msg = quote("No tienes permisos para gestionar emisores/receptores")
            return RedirectResponse(url=f"/?error={msg}", status_code=302)

    # Si no hay usuario en sesión -> login
    if not request.session.get("user"):
        return RedirectResponse(url=f"/login?next={path}", status_code=302)

    return await call_next(request)


app.add_middleware(
    SessionMiddleware,
    secret_key="dev-secret-cambia-esto",
    session_cookie="facturador_session",
)



def format_cop(value):
    if value is None:
        value = 0
    dec = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    formatted = f"{dec:,.2f}"
    formatted = formatted.replace(",", "X").replace(".", ",").replace("X", ".")
    return f"${formatted}"

def currency_cop(value):
    return format_cop(value)


def parse_cop_to_decimal(value):
    if value is None:
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip()
    if text == "":
        return Decimal("0")
    text = text.replace("$", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "")
        text = text.replace(",", ".")
    elif text.count(".") > 1:
        text = text.replace(".", "")
    elif "." in text:
        left, right = text.split(".", 1)
        if len(right) > 2:
            text = left + right
    text = "".join(ch for ch in text if ch.isdigit() or ch == ".")
    if text in ("", "."):
        return Decimal("0")
    try:
        return Decimal(text)
    except Exception:
        return Decimal("0")


# Backwards-compatible alias
def parse_cop(value):
    return parse_cop_to_decimal(value)


def to_sqlite_number(value):
    if isinstance(value, Decimal):
        return str(value)
    return value


def digits_only(value):
    if value is None:
        return ""
    return "".join(ch for ch in str(value) if ch.isdigit())


def parse_checkbox_value(value):
    if value is None:
        return 0
    return 1 if str(value).strip().lower() in ("1", "true", "on", "yes") else 0


def format_id_co(value):
    digits = digits_only(value)
    if digits == "":
        return ""
    rev = digits[::-1]
    chunks = [rev[i : i + 3] for i in range(0, len(rev), 3)]
    return ".".join(chunk[::-1] for chunk in chunks[::-1])


def nit_dv(value):
    digits = digits_only(value)
    if not digits:
        return ""
    weights = [3, 7, 13, 17, 19, 23, 29, 37, 41, 43, 47, 53, 59, 67, 71]
    total = 0
    idx = 0
    for ch in reversed(digits):
        total += int(ch) * weights[idx]
        idx += 1
    mod = total % 11
    return str(mod if mod in (0, 1) else 11 - mod)


def format_tipo_display(value):
    if value is None:
        return ""
    if value == "FACTURA":
        return "Factura de venta"
    if value == "CUENTA_COBRO":
        return "Cuenta de cobro"
    return value


def _num_to_words_es(number):
    units = [
        "", "UNO", "DOS", "TRES", "CUATRO", "CINCO",
        "SEIS", "SIETE", "OCHO", "NUEVE",
    ]
    specials = {
        10: "DIEZ", 11: "ONCE", 12: "DOCE", 13: "TRECE", 14: "CATORCE",
        15: "QUINCE", 16: "DIECISEIS", 17: "DIECISIETE", 18: "DIECIOCHO",
        19: "DIECINUEVE",
    }
    tens_words = {
        2: "VEINTE", 3: "TREINTA", 4: "CUARENTA", 5: "CINCUENTA",
        6: "SESENTA", 7: "SETENTA", 8: "OCHENTA", 9: "NOVENTA",
    }
    hundreds_words = {
        1: "CIENTO", 2: "DOSCIENTOS", 3: "TRESCIENTOS", 4: "CUATROCIENTOS",
        5: "QUINIENTOS", 6: "SEISCIENTOS", 7: "SETECIENTOS",
        8: "OCHOCIENTOS", 9: "NOVECIENTOS",
    }

    if number == 0:
        return "CERO"
    if number < 10:
        return units[number]
    if number in specials:
        return specials[number]
    if number < 30:
        return "VEINTI" + units[number - 20]
    if number < 100:
        ten = number // 10
        unit = number % 10
        if unit == 0:
            return tens_words[ten]
        return f"{tens_words[ten]} Y {units[unit]}"
    if number == 100:
        return "CIEN"
    if number < 1000:
        hundred = number // 100
        rest = number % 100
        if rest == 0:
            return hundreds_words[hundred]
        return f"{hundreds_words[hundred]} {_num_to_words_es(rest)}"
    if number < 1_000_000:
        thousand = number // 1000
        rest = number % 1000
        prefix = "MIL" if thousand == 1 else f"{_num_to_words_es(thousand)} MIL"
        if rest == 0:
            return prefix
        return f"{prefix} {_num_to_words_es(rest)}"
    if number < 1_000_000_000:
        million = number // 1_000_000
        rest = number % 1_000_000
        prefix = "UN MILLON" if million == 1 else f"{_num_to_words_es(million)} MILLONES"
        if rest == 0:
            return prefix
        return f"{prefix} {_num_to_words_es(rest)}"
    billion = number // 1_000_000_000
    rest = number % 1_000_000_000
    prefix = "MIL MILLONES" if billion == 1 else f"{_num_to_words_es(billion)} MIL MILLONES"
    if rest == 0:
        return prefix
    return f"{prefix} {_num_to_words_es(rest)}"


def num_to_words_cop(value):
    if value is None:
        value = 0
    dec = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    entero = int(dec)
    centavos = int((dec - Decimal(entero)) * 100)
    words = _num_to_words_es(entero)
    if words.endswith("UNO"):
        words = words[:-1]
    moneda = "PESO" if entero == 1 else "PESOS"
    return f"{words} {moneda} CON {centavos:02d}/100"


def build_document_pdf(doc, lines):
    if not isinstance(doc, dict):
        doc = dict(doc)
    lines = [dict(line) if not isinstance(line, dict) else line for line in lines]
    buffer = BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("Title", parent=styles["Heading1"], alignment=1)
    center = ParagraphStyle("Center", parent=styles["Normal"], alignment=1, leading=16)
    normal = ParagraphStyle("Normal", parent=styles["Normal"], fontSize=10, leading=14)
    label = ParagraphStyle("Label", parent=styles["Normal"], fontSize=9, textColor=colors.grey)
    note = ParagraphStyle("Note", parent=styles["Normal"], fontSize=8, leading=11, textColor=colors.grey)

    def make_block(title_text, values):
        parts = [Paragraph(f"<b>{title_text}</b>", label)]
        for val in values:
            if val:
                parts.append(Paragraph(str(val), normal))
        return parts

    def format_id_line(tipo, identificacion):
        if not identificacion:
            return ""
        tipo_clean = (tipo or "").strip().upper()
        base = format_id_co(identificacion)
        if tipo_clean == "NIT":
            dv = nit_dv(identificacion)
            if dv:
                return f"{tipo_clean} {base}-{dv}"
        return f"{tipo_clean} {base}".strip()

    story = []
    numero = f"{doc['prefijo']}-{doc['numero']:06d}"

    if doc["tipo"] == "CUENTA_COBRO":
        story.append(Paragraph("CUENTA DE COBRO", title))
        story.append(Paragraph(f"No. {numero}", center))
        story.append(Paragraph(f"Fecha: {doc['fecha']}", center))
        story.append(Spacer(1, 12))
        for line in [
            doc["receptor_nombre"],
            format_id_line(doc.get("receptor_tipo_identificacion"), doc.get("receptor_identificacion")),
            doc.get("receptor_ciudad"),
            doc.get("receptor_direccion"),
            doc.get("receptor_email"),
            doc.get("receptor_telefono"),
        ]:
            if line:
                story.append(Paragraph(str(line), center))
        story.append(Spacer(1, 12))
        story.append(Paragraph("<b>DEBE A</b>", center))
        story.append(Spacer(1, 12))
        for line in [
            doc["emisor_nombre"],
            format_id_line(doc.get("emisor_tipo_identificacion"), doc.get("emisor_identificacion")),
            doc.get("emisor_ciudad"),
            doc.get("emisor_direccion"),
            doc.get("emisor_email"),
            doc.get("emisor_telefono"),
        ]:
            if line:
                story.append(Paragraph(str(line), center))
        story.append(Spacer(1, 16))
    else:
        story.append(Paragraph("FACTURA DE VENTA", title))
        story.append(Paragraph(numero, center))
        story.append(Paragraph(f"Fecha: {doc['fecha']}", center))
        story.append(Spacer(1, 20))
        emisor_block = make_block(
            "Emisor",
            [
                doc["emisor_nombre"],
                format_id_line(doc.get("emisor_tipo_identificacion"), doc.get("emisor_identificacion")),
                doc.get("emisor_ciudad"),
                doc.get("emisor_direccion"),
                doc.get("emisor_email"),
                doc.get("emisor_telefono"),
            ],
        )
        receptor_block = make_block(
            "Cliente",
            [
                doc["receptor_nombre"],
                format_id_line(doc.get("receptor_tipo_identificacion"), doc.get("receptor_identificacion")),
                doc.get("receptor_ciudad"),
                doc.get("receptor_direccion"),
                doc.get("receptor_email"),
                doc.get("receptor_telefono"),
            ],
        )
        info_table = Table([[emisor_block, receptor_block]], colWidths=[pdf.width / 2] * 2)
        info_table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        story.append(info_table)
        story.append(Spacer(1, 12))

    table_data = [
        ["Item", "Descripcion", "Cantidad", "Precio", "IVA", "Total"]
    ]
    for line in lines:
        table_data.append(
            [
                line.get("item_codigo") or line.get("item_descripcion") or line.get("descripcion") or "",
                line.get("descripcion") or "",
                f"{line['cantidad']:.2f}",
                currency_cop(line["precio_unitario"]),
                currency_cop(line["iva"]),
                currency_cop(line["total"]),
            ]
        )
    table = Table(table_data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f4f8")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
                ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 10))

    totals = Table(
        [
            ["Subtotal", currency_cop(doc["subtotal"])],
            ["IVA", currency_cop(doc["iva"])],
            ["Total", currency_cop(doc["total"])],
        ],
        colWidths=[pdf.width * 0.7, pdf.width * 0.3],
    )
    totals.setStyle(
        TableStyle(
            [
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(totals)
    story.append(Spacer(1, 6))
    story.append(
        Paragraph(
            f"<b>Valor a pagar en letras:</b> {num_to_words_cop(doc['total'])}",
            normal,
        )
    )
    if doc.get("observaciones"):
        story.append(Spacer(1, 8))
        story.append(Paragraph("<b>Observaciones:</b>", normal))
        for line in str(doc.get("observaciones")).splitlines():
            story.append(Paragraph(line, normal))
    if doc.get("include_tax_note_383") and doc.get("tipo") == "CUENTA_COBRO":
        story.append(Spacer(1, 6))
        story.append(Paragraph(TAX_NOTE_383, note))

    pdf.build(story)
    buffer.seek(0)
    return buffer


templates.env.filters["currency_cop"] = currency_cop
templates.env.filters["num_to_words_cop"] = num_to_words_cop
templates.env.filters["format_id_co"] = format_id_co
templates.env.filters["nit_dv"] = nit_dv
templates.env.filters["format_tipo_display"] = format_tipo_display
TAX_NOTE_383 = (
    "Nota tributaria: El suscrito manifiesta bajo la gravedad de juramento que los "
    "ingresos objeto de este cobro corresponden a rentas de trabajo no provenientes "
    "de una relación laboral o legal y reglamentaria y solicita que la retención en "
    "la fuente se determine conforme al artículo 383 del E.T., en los términos del "
    "parágrafo 4 del artículo 1.2.4.1.17 del D. 1625 de 2016, por no solicitar la "
    "aplicación de costos y deducciones. Así mismo, para efectos de la aplicación "
    "del numeral 10 del artículo 206 del E.T., declara que no tomará costos ni "
    "deducciones asociadas a dichas rentas."
)
templates.env.globals["tax_note_383"] = TAX_NOTE_383
def format_percent(value):
    if value is None:
        return "0%"
    try:
        val = float(value)
    except (TypeError, ValueError):
        return "0%"
    if val <= 1:
        val = val * 100
    if abs(val - round(val)) < 1e-9:
        return f"{int(round(val))}%"
    return f"{val:.2f}%"


templates.env.filters["format_percent"] = format_percent

DB_PATH = os.path.join(os.path.dirname(__file__), "facturador.db")
DATA_PATH = os.path.join(os.path.dirname(__file__), "data", "municipios_co.json")
ITEM_HAS_NOMBRE = False


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def ensure_columns(conn, table_name, columns):
    cur = conn.cursor()
    cur.execute(f"PRAGMA table_info({table_name})")
    existing = {row["name"] for row in cur.fetchall()}
    for name, col_def in columns.items():
        if name not in existing:
            cur.execute(f"ALTER TABLE {table_name} ADD COLUMN {col_def}")
    conn.commit()


def get_or_create_group(conn, name):
    cur = conn.cursor()
    cur.execute("SELECT id FROM grupos WHERE lower(nombre) = lower(?)", (name,))
    row = cur.fetchone()
    if row:
        return row["id"]
    cur.execute("INSERT INTO grupos (nombre) VALUES (?)", (name,))
    conn.commit()
    return cur.lastrowid


def migrate_item_groups(conn):
    cur = conn.cursor()
    cur.execute("PRAGMA table_info(items)")
    cols = {row["name"] for row in cur.fetchall()}
    if "grupo" not in cols:
        return

    cur.execute(
        "SELECT DISTINCT grupo FROM items WHERE grupo IS NOT NULL AND trim(grupo) != ''"
    )
    groups = [row["grupo"] for row in cur.fetchall()]
    for name in groups:
        group_id = get_or_create_group(conn, name)
        cur.execute(
            """
            UPDATE items
            SET grupo_id = ?
            WHERE grupo_id IS NULL AND lower(grupo) = lower(?)
            """,
            (group_id, name),
        )
    conn.commit()


def migrate_item_codes(conn):
    cur = conn.cursor()
    cur.execute("PRAGMA table_info(items)")
    cols = {row["name"] for row in cur.fetchall()}
    if "codigo" not in cols or "nombre" not in cols:
        return

    cur.execute(
        "UPDATE items SET codigo = nombre WHERE (codigo IS NULL OR trim(codigo) = '')"
    )
    conn.commit()


def ensure_unique_item_codes(conn):
    cur = conn.cursor()
    cur.execute("PRAGMA table_info(items)")
    cols = {row["name"] for row in cur.fetchall()}
    has_nombre = "nombre" in cols

    cur.execute(
        """
        SELECT lower(codigo) AS key, COUNT(*) AS c
        FROM items
        WHERE codigo IS NOT NULL AND trim(codigo) != ''
        GROUP BY lower(codigo)
        HAVING COUNT(*) > 1
        """
    )
    duplicates = [row["key"] for row in cur.fetchall()]
    if not duplicates:
        return

    for key in duplicates:
        cur.execute(
            """
            SELECT id, codigo
            FROM items
            WHERE lower(codigo) = lower(?)
            ORDER BY id
            """,
            (key,),
        )
        rows = cur.fetchall()
        for row in rows[1:]:
            new_code = f"{row['codigo']}-{row['id']}"
            if has_nombre:
                cur.execute(
                    "UPDATE items SET codigo = ?, nombre = ? WHERE id = ?",
                    (new_code, new_code, row["id"]),
                )
            else:
                cur.execute(
                    "UPDATE items SET codigo = ? WHERE id = ?",
                    (new_code, row["id"]),
                )
    conn.commit()


def init_db():
    global ITEM_HAS_NOMBRE
    conn = get_db()
    cur = conn.cursor()

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS emisores (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT NOT NULL,
            identificacion TEXT NOT NULL,
            ciudad TEXT NOT NULL DEFAULT '',
            departamento TEXT NOT NULL DEFAULT '',
            pais TEXT NOT NULL DEFAULT 'Colombia',
            direccion TEXT,
            telefono TEXT,
            email TEXT
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS receptores (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT NOT NULL,
            identificacion TEXT NOT NULL,
            ciudad TEXT NOT NULL DEFAULT '',
            departamento TEXT NOT NULL DEFAULT '',
            pais TEXT NOT NULL DEFAULT 'Colombia',
            direccion TEXT,
            telefono TEXT,
            email TEXT
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS grupos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT NOT NULL COLLATE NOCASE UNIQUE
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS items (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            codigo TEXT NOT NULL,
            descripcion TEXT,
            costo REAL NOT NULL DEFAULT 0.0,
            precio REAL NOT NULL,
            iva_rate REAL NOT NULL DEFAULT 0.0,
            grupo_id INTEGER,
            FOREIGN KEY (grupo_id) REFERENCES grupos(id)
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS consecutivos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            emisor_id INTEGER NOT NULL,
            tipo TEXT NOT NULL,
            prefijo TEXT NOT NULL,
            ultimo_numero INTEGER NOT NULL DEFAULT 0,
            UNIQUE (emisor_id, tipo, prefijo),
            FOREIGN KEY (emisor_id) REFERENCES emisores(id)
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS documentos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            emisor_id INTEGER NOT NULL,
            receptor_id INTEGER NOT NULL,
            tipo TEXT NOT NULL,
            prefijo TEXT NOT NULL,
            numero INTEGER,
            fecha TEXT NOT NULL,
            subtotal REAL NOT NULL,
            iva REAL NOT NULL,
            total REAL NOT NULL,
            observaciones TEXT,
            include_tax_note_383 INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            estado TEXT NOT NULL DEFAULT 'EMITIDO',
            FOREIGN KEY (emisor_id) REFERENCES emisores(id),
            FOREIGN KEY (receptor_id) REFERENCES receptores(id)
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS documento_detalle (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            documento_id INTEGER NOT NULL,
            item_id INTEGER,
            descripcion TEXT NOT NULL,
            cantidad REAL NOT NULL,
            precio_unitario REAL NOT NULL,
            iva_rate REAL NOT NULL,
            subtotal REAL NOT NULL,
            iva REAL NOT NULL,
            total REAL NOT NULL,
            FOREIGN KEY (documento_id) REFERENCES documentos(id),
            FOREIGN KEY (item_id) REFERENCES items(id)
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS documentos_borrador (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            emisor_id INTEGER NOT NULL,
            receptor_id INTEGER NOT NULL,
            tipo TEXT NOT NULL,
            prefijo TEXT NOT NULL,
            fecha TEXT NOT NULL,
            subtotal REAL NOT NULL,
            iva REAL NOT NULL,
            total REAL NOT NULL,
            created_at TEXT NOT NULL,
            FOREIGN KEY (emisor_id) REFERENCES emisores(id),
            FOREIGN KEY (receptor_id) REFERENCES receptores(id)
        )
        """
    )

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS documento_borrador_detalle (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            borrador_id INTEGER NOT NULL,
            item_id INTEGER,
            descripcion TEXT NOT NULL,
            cantidad REAL NOT NULL,
            precio_unitario REAL NOT NULL,
            iva_rate REAL NOT NULL,
            subtotal REAL NOT NULL,
            iva REAL NOT NULL,
            total REAL NOT NULL,
            FOREIGN KEY (borrador_id) REFERENCES documentos_borrador(id),
            FOREIGN KEY (item_id) REFERENCES items(id)
        )
        """
    )

    conn.commit()

    migrate_documentos_estado(conn)
    ensure_columns(
        conn,
        "documentos",
        {"include_tax_note_383": "include_tax_note_383 INTEGER NOT NULL DEFAULT 0"},
    )

    ensure_columns(
        conn,
        "emisores",
        {
            "ciudad": "ciudad TEXT NOT NULL DEFAULT ''",
            "departamento": "departamento TEXT NOT NULL DEFAULT ''",
            "pais": "pais TEXT NOT NULL DEFAULT 'Colombia'",
            "tipo_identificacion": "tipo_identificacion TEXT NOT NULL DEFAULT 'NIT'",
        },
    )
    ensure_columns(
        conn,
        "receptores",
        {
            "ciudad": "ciudad TEXT NOT NULL DEFAULT ''",
            "departamento": "departamento TEXT NOT NULL DEFAULT ''",
            "pais": "pais TEXT NOT NULL DEFAULT 'Colombia'",
            "tipo_identificacion": "tipo_identificacion TEXT NOT NULL DEFAULT 'CC'",
        },
    )
    ensure_columns(
        conn,
        "items",
        {
            "costo": "costo REAL NOT NULL DEFAULT 0.0",
            "grupo_id": "grupo_id INTEGER",
            "codigo": "codigo TEXT",
        },
    )

    cur.execute("PRAGMA table_info(items)")
    ITEM_HAS_NOMBRE = "nombre" in {row["name"] for row in cur.fetchall()}

    migrate_item_groups(conn)
    migrate_item_codes(conn)
    ensure_unique_item_codes(conn)
    cur.execute(
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_items_codigo_lower ON items(lower(codigo))"
    )

    cur.execute(
        "UPDATE emisores SET tipo_identificacion = 'NIT' WHERE tipo_identificacion IS NULL OR tipo_identificacion = ''"
    )
    cur.execute(
        "UPDATE receptores SET tipo_identificacion = 'CC' WHERE tipo_identificacion IS NULL OR tipo_identificacion = ''"
    )

    conn.commit()
    conn.close()


def migrate_documentos_estado(conn):
    cur = conn.cursor()
    cur.execute("PRAGMA table_info(documentos)")
    cols = cur.fetchall()
    col_names = {row["name"] for row in cols}
    numero_notnull = None
    for row in cols:
        if row["name"] == "numero":
            numero_notnull = row["notnull"]

    needs_migration = (
        "estado" not in col_names
        or numero_notnull == 1
        or "observaciones" not in col_names
        or "include_tax_note_383" not in col_names
    )
    if not needs_migration:
        cur.execute(
            """
            CREATE UNIQUE INDEX IF NOT EXISTS idx_documentos_consecutivo
            ON documentos(emisor_id, tipo, prefijo, numero)
            WHERE estado = 'EMITIDO' AND numero IS NOT NULL
            """
        )
        if "observaciones" not in col_names:
            cur.execute("ALTER TABLE documentos ADD COLUMN observaciones TEXT")
        if "include_tax_note_383" not in col_names:
            cur.execute(
                "ALTER TABLE documentos ADD COLUMN include_tax_note_383 INTEGER NOT NULL DEFAULT 0"
            )
        conn.commit()
        return

    cur.execute("PRAGMA foreign_keys = OFF")
    cur.execute(
        """
        CREATE TABLE documentos_new (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            emisor_id INTEGER NOT NULL,
            receptor_id INTEGER NOT NULL,
            tipo TEXT NOT NULL,
            prefijo TEXT NOT NULL,
            numero INTEGER,
            fecha TEXT NOT NULL,
            subtotal REAL NOT NULL,
            iva REAL NOT NULL,
            total REAL NOT NULL,
            observaciones TEXT,
            include_tax_note_383 INTEGER NOT NULL DEFAULT 0,
            created_at TEXT NOT NULL,
            estado TEXT NOT NULL DEFAULT 'EMITIDO',
            FOREIGN KEY (emisor_id) REFERENCES emisores(id),
            FOREIGN KEY (receptor_id) REFERENCES receptores(id)
        )
        """
    )
    cur.execute(
        """
        INSERT INTO documentos_new (
            id, emisor_id, receptor_id, tipo, prefijo, numero, fecha,
            subtotal, iva, total, observaciones, include_tax_note_383, created_at, estado
        )
        SELECT id, emisor_id, receptor_id, tipo, prefijo, numero, fecha,
               subtotal, iva, total, NULL, 0, created_at, 'EMITIDO'
        FROM documentos
        """
    )
    cur.execute("DROP TABLE documentos")
    cur.execute("ALTER TABLE documentos_new RENAME TO documentos")
    cur.execute(
        """
        CREATE UNIQUE INDEX IF NOT EXISTS idx_documentos_consecutivo
        ON documentos(emisor_id, tipo, prefijo, numero)
        WHERE estado = 'EMITIDO' AND numero IS NOT NULL
        """
    )
    cur.execute("PRAGMA foreign_keys = ON")
    conn.commit()


def get_consecutivo(conn, emisor_id, tipo, prefijo):
    cur = conn.cursor()
    cur.execute(
        """
        SELECT id, ultimo_numero
        FROM consecutivos
        WHERE emisor_id = ? AND tipo = ? AND prefijo = ?
        """,
        (emisor_id, tipo, prefijo),
    )
    row = cur.fetchone()
    if row is None:
        cur.execute(
            """
            INSERT INTO consecutivos (emisor_id, tipo, prefijo, ultimo_numero)
            VALUES (?, ?, ?, 0)
            """,
            (emisor_id, tipo, prefijo),
        )
        conn.commit()
        consec_id = cur.lastrowid
        ultimo = 0
    else:
        consec_id = row["id"]
        ultimo = row["ultimo_numero"]

    siguiente = int(ultimo) + 1
    cur.execute(
        "UPDATE consecutivos SET ultimo_numero = ? WHERE id = ?",
        (siguiente, consec_id),
    )
    conn.commit()
    return siguiente


def get_next_consecutivo(conn, emisor_id, tipo, prefijo):
    cur = conn.cursor()
    cur.execute(
        """
        SELECT ultimo_numero
        FROM consecutivos
        WHERE emisor_id = ? AND tipo = ? AND prefijo = ?
        """,
        (emisor_id, tipo, prefijo),
    )
    row = cur.fetchone()
    last_num = int(row["ultimo_numero"]) if row else 0
    return last_num + 1


def calculate_totals(lines):
    subtotal = 0.0
    iva = 0.0
    total = 0.0

    for line in lines:
        line_sub = float(line["cantidad"]) * float(line["precio_unitario"])
        line_iva = line_sub * float(line["iva_rate"])
        line_total = line_sub + line_iva
        line["subtotal"] = line_sub
        line["iva"] = line_iva
        line["total"] = line_total
        subtotal += line_sub
        iva += line_iva
        total += line_total

    return subtotal, iva, total


def build_lines_from_form(form, items_rows):
    lines = []
    indices = set()
    for key in form.keys():
        if not key.startswith("item_id_"):
            continue
        try:
            indices.add(int(key.split("_", 2)[2]))
        except (IndexError, ValueError):
            continue

    for idx in sorted(indices):
        item_id = form.get(f"item_id_{idx}")
        cantidad = form.get(f"cantidad_{idx}")
        precio = form.get(f"precio_{idx}")
        iva_percent = form.get(f"iva_{idx}")
        if not item_id or not cantidad:
            continue
        try:
            cantidad_val = float(cantidad)
        except ValueError:
            continue
        if cantidad_val <= 0:
            continue
        lines.append(
            {
                "item_id": int(item_id),
                "cantidad": cantidad_val,
                "precio_unitario": precio,
                "iva_percent": iva_percent,
            }
        )

    doc_lines = []
    for line in lines:
        item = items_rows.get(line["item_id"])
        if not item:
            continue
        descripcion = item["descripcion"] or item["codigo"]
        try:
            precio_unitario = parse_cop(line["precio_unitario"] or item["precio"])
        except ValueError:
            precio_unitario = float(item["precio"])
        try:
            iva_percent = float(line["iva_percent"])
        except (TypeError, ValueError):
            iva_percent = float(item["iva_rate"]) * 100
        iva_rate = iva_percent / 100.0
        doc_lines.append(
            {
                "item_id": item["id"],
                "descripcion": descripcion,
                "cantidad": line["cantidad"],
                "precio_unitario": precio_unitario,
                "iva_rate": iva_rate,
            }
        )
    return doc_lines


def parse_document_filters(qp, default_estado=None):
    def parse_int_optional(value):
        if value is None:
            return None
        value = str(value).strip()
        if value == "":
            return None
        try:
            return int(value)
        except ValueError:
            return None

    emisor_id = parse_int_optional(qp.get("emisor_id"))
    receptor_id = parse_int_optional(qp.get("receptor_id"))
    numero = parse_int_optional(qp.get("numero"))
    tipo_documento = (qp.get("tipo_documento") or "").strip() or None
    prefijo = (qp.get("prefijo") or "").strip() or None
    fecha_desde = (qp.get("fecha_desde") or "").strip() or None
    fecha_hasta = (qp.get("fecha_hasta") or "").strip() or None
    estado = (qp.get("estado") or "").strip() or None

    if estado is None and default_estado is not None:
        estado = default_estado
    if estado == "TODOS":
        estado = None

    where = []
    params = []
    if emisor_id:
        where.append("d.emisor_id = ?")
        params.append(emisor_id)
    if receptor_id:
        where.append("d.receptor_id = ?")
        params.append(receptor_id)
    if tipo_documento:
        where.append("d.tipo = ?")
        params.append(tipo_documento)
    if prefijo:
        where.append("d.prefijo = ?")
        params.append(prefijo)
    if estado:
        where.append("d.estado = ?")
        params.append(estado)
    if numero:
        where.append("d.numero = ?")
        params.append(numero)
    if fecha_desde:
        where.append("d.fecha >= ?")
        params.append(fecha_desde)
    if fecha_hasta:
        where.append("d.fecha <= ?")
        params.append(fecha_hasta)

    filters = {
        "emisor_id": emisor_id,
        "receptor_id": receptor_id,
        "tipo_documento": tipo_documento,
        "numero": numero,
        "prefijo": prefijo,
        "fecha_desde": fecha_desde,
        "fecha_hasta": fecha_hasta,
        "estado": estado or "TODOS",
    }
    return where, params, filters


def build_filters_query_string(filters):
    parts = []
    for key, value in filters.items():
        if value is None:
            continue
        if isinstance(value, str) and value.strip() == "":
            continue
        if value == "TODOS":
            continue
        parts.append((key, str(value)))
    if not parts:
        return ""
    return "?" + "&".join(f"{k}={quote(v)}" for k, v in parts)


def fetch_documents(conn, where, params, limit=None):
    query = """
        SELECT d.*, e.nombre AS emisor_nombre, r.nombre AS receptor_nombre
        FROM documentos d
        JOIN emisores e ON e.id = d.emisor_id
        JOIN receptores r ON r.id = d.receptor_id
    """
    if where:
        query += " WHERE " + " AND ".join(where)
    query += " ORDER BY d.fecha DESC, d.id DESC"
    if limit:
        query += f" LIMIT {int(limit)}"
    cur = conn.cursor()
    cur.execute(query, params)
    return cur.fetchall()


def fetch_documents_sum(conn, where, params):
    sum_query = "SELECT COALESCE(SUM(d.total), 0) AS total_sum FROM documentos d"
    if where:
        sum_query += " WHERE " + " AND ".join(where)
    cur = conn.cursor()
    cur.execute(sum_query, params)
    row = cur.fetchone()
    return row["total_sum"] if row else 0


def format_filters_text(filters):
    parts = []
    labels = {
        "emisor_id": "Emisor",
        "receptor_id": "Cliente",
        "tipo_documento": "Tipo",
        "estado": "Estado",
        "prefijo": "Prefijo",
        "numero": "Numero",
        "fecha_desde": "Fecha desde",
        "fecha_hasta": "Fecha hasta",
    }
    for key, label in labels.items():
        val = filters.get(key)
        if not val or val == "TODOS":
            continue
        parts.append(f"{label}: {val}")
    return " | ".join(parts) if parts else "Sin filtros"


def format_id_with_type(tipo, identificacion):
    if not identificacion:
        return ""
    tipo_clean = (tipo or "").strip().upper()
    base = format_id_co(identificacion)
    if tipo_clean == "NIT":
        dv = nit_dv(identificacion)
        if dv:
            return f"{tipo_clean} {base}-{dv}"
    return f"{tipo_clean} {base}".strip()


def fetch_report_lines(conn, where, params):
    query = """
        SELECT d.id AS documento_id, d.fecha, d.tipo, d.prefijo, d.numero,
               r.nombre AS receptor_nombre,
               dd.descripcion, dd.cantidad, dd.precio_unitario, dd.iva_rate,
               dd.subtotal, dd.iva, dd.total
        FROM documentos d
        JOIN documento_detalle dd ON dd.documento_id = d.id
        LEFT JOIN receptores r ON r.id = d.receptor_id
    """
    if where:
        query += " WHERE " + " AND ".join(where)
    query += " ORDER BY d.fecha DESC, d.id DESC, dd.id ASC"
    cur = conn.cursor()
    cur.execute(query, params)
    return cur.fetchall()


def build_ventas_lines(lines_raw, mode: str, report: str):
    if report not in ("ventas_cliente", "ventas_producto"):
        report = "ventas_cliente"

    lines = []
    total_qty = 0.0
    total_valor = 0.0
    current_primary = None
    current_secondary = None
    sub_qty = 0.0
    sub_total = 0.0
    prod_qty = 0.0
    prod_total = 0.0

    def flush_secondary():
        if current_primary is None or current_secondary is None:
            return
        lines.append(
            {
                "is_subtotal": False,
                "primary": current_primary,
                "secondary": current_secondary,
                "cantidad": prod_qty,
                "valor_total": prod_total,
            }
        )

    def flush_subtotal():
        if current_primary is None:
            return
        lines.append(
            {
                "is_subtotal": True,
                "primary": current_primary,
                "secondary": "",
                "cantidad": sub_qty,
                "valor_total": sub_total,
            }
        )

    def to_keys(line):
        cliente = line["receptor_nombre"] or "-"
        descripcion = line["descripcion"] or ""
        if report == "ventas_producto":
            return descripcion, cliente
        return cliente, descripcion

    lines_sorted = sorted(
        lines_raw,
        key=lambda line: (
            (to_keys(line)[0] or "").lower(),
            (to_keys(line)[1] or "").lower(),
            line["fecha"],
            line["documento_id"],
        ),
    )
    for line in lines_sorted:
        primary, secondary = to_keys(line)
        if current_primary is None:
            current_primary = primary
            current_secondary = secondary
        if primary != current_primary:
            if mode == "resumen":
                flush_secondary()
            flush_subtotal()
            current_primary = primary
            current_secondary = secondary
            sub_qty = 0.0
            sub_total = 0.0
            prod_qty = 0.0
            prod_total = 0.0
        if mode == "resumen" and secondary != current_secondary:
            flush_secondary()
            current_secondary = secondary
            prod_qty = 0.0
            prod_total = 0.0

        qty = float(line["cantidad"] or 0)
        subtotal = float(line["subtotal"] or 0)
        total_qty += qty
        total_valor += subtotal
        sub_qty += qty
        sub_total += subtotal
        prod_qty += qty
        prod_total += subtotal

        if mode == "detalle":
            numero = "-"
            if line["numero"] is not None:
                numero = f"{line['prefijo']}-{str(line['numero']).zfill(6)}"
            lines.append(
                {
                    "is_subtotal": False,
                    "primary": primary,
                    "secondary": secondary,
                    "documento": numero,
                    "cantidad": qty,
                    "valor_unitario": float(line["precio_unitario"] or 0),
                    "valor_total": subtotal,
                }
            )

    if mode == "resumen":
        flush_secondary()
    flush_subtotal()
    totals = {"cantidad": total_qty, "total": total_valor}
    return lines, totals




def create_document(
    conn, emisor_id, receptor_id, tipo, prefijo, lines, fecha, observaciones, include_tax_note_383
):
    subtotal, iva, total = calculate_totals(lines)

    numero = get_consecutivo(conn, emisor_id, tipo, prefijo)
    created_at = datetime.datetime.now().isoformat(timespec="seconds")

    cur = conn.cursor()
    cur.execute(
        """
        INSERT INTO documentos (
            emisor_id, receptor_id, tipo, prefijo, numero, fecha,
            subtotal, iva, total, observaciones, include_tax_note_383, created_at, estado
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            emisor_id,
            receptor_id,
            tipo,
            prefijo,
            numero,
            fecha,
            to_sqlite_number(subtotal),
            to_sqlite_number(iva),
            to_sqlite_number(total),
            observaciones,
            include_tax_note_383,
            created_at,
            "EMITIDO",
        ),
    )
    doc_id = cur.lastrowid

    for line in lines:
        cur.execute(
            """
            INSERT INTO documento_detalle (
                documento_id, item_id, descripcion, cantidad, precio_unitario,
                iva_rate, subtotal, iva, total
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                doc_id,
                line.get("item_id"),
                line["descripcion"],
                to_sqlite_number(line["cantidad"]),
                to_sqlite_number(line["precio_unitario"]),
                to_sqlite_number(line["iva_rate"]),
                to_sqlite_number(line["subtotal"]),
                to_sqlite_number(line["iva"]),
                to_sqlite_number(line["total"]),
            ),
        )

    conn.commit()
    return doc_id, numero


def create_draft_document(
    conn, emisor_id, receptor_id, tipo, prefijo, lines, fecha, observaciones, numero, include_tax_note_383
):
    subtotal, iva, total = calculate_totals(lines)
    created_at = datetime.datetime.now().isoformat(timespec="seconds")

    cur = conn.cursor()
    cur.execute(
        """
        INSERT INTO documentos (
            emisor_id, receptor_id, tipo, prefijo, numero, fecha,
            subtotal, iva, total, observaciones, include_tax_note_383, created_at, estado
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            emisor_id,
            receptor_id,
            tipo,
            prefijo,
            numero,
            fecha,
            to_sqlite_number(subtotal),
            to_sqlite_number(iva),
            to_sqlite_number(total),
            observaciones,
            include_tax_note_383,
            created_at,
            "BORRADOR",
        ),
    )
    doc_id = cur.lastrowid

    for line in lines:
        cur.execute(
            """
            INSERT INTO documento_detalle (
                documento_id, item_id, descripcion, cantidad, precio_unitario,
                iva_rate, subtotal, iva, total
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                doc_id,
                line.get("item_id"),
                line["descripcion"],
                to_sqlite_number(line["cantidad"]),
                to_sqlite_number(line["precio_unitario"]),
                to_sqlite_number(line["iva_rate"]),
                to_sqlite_number(line["subtotal"]),
                to_sqlite_number(line["iva"]),
                to_sqlite_number(line["total"]),
            ),
        )

    conn.commit()
    return doc_id


def fetch_document(conn, doc_id):
    cur = conn.cursor()
    cur.execute(
        """
        SELECT d.*, e.nombre AS emisor_nombre, r.nombre AS receptor_nombre,
               e.identificacion AS emisor_identificacion, r.identificacion AS receptor_identificacion,
               e.tipo_identificacion AS emisor_tipo_identificacion,
               r.tipo_identificacion AS receptor_tipo_identificacion,
               e.ciudad AS emisor_ciudad, e.departamento AS emisor_departamento,
               r.ciudad AS receptor_ciudad, r.departamento AS receptor_departamento,
               e.direccion AS emisor_direccion, e.email AS emisor_email, e.telefono AS emisor_telefono,
               r.direccion AS receptor_direccion, r.email AS receptor_email, r.telefono AS receptor_telefono
        FROM documentos d
        JOIN emisores e ON e.id = d.emisor_id
        JOIN receptores r ON r.id = d.receptor_id
        WHERE d.id = ?
        """,
        (doc_id,),
    )
    doc = cur.fetchone()
    cur.execute(
        """
        SELECT dd.*, i.descripcion AS item_descripcion, i.codigo AS item_codigo
        FROM documento_detalle dd
        LEFT JOIN items i ON i.id = dd.item_id
        WHERE dd.documento_id = ?
        ORDER BY dd.id
        """,
        (doc_id,),
    )
    lines = cur.fetchall()
    return doc, lines


@app.get("/", response_class=HTMLResponse)
def index(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})


@app.get("/data/cities", response_class=JSONResponse)
def cities_data():
    if not os.path.exists(DATA_PATH):
        return []
    with open(DATA_PATH, "r", encoding="utf-8") as handle:
        return json.load(handle)


@app.get("/consecutivo/preview", response_class=JSONResponse)
def consecutivo_preview(emisor_id: int | None = None, tipo: str | None = None, prefijo: str | None = None):
    if not emisor_id or not tipo:
        return {"next_numero": None}
    pref = (prefijo or "").strip()
    if not pref:
        pref = "FV" if tipo == "FACTURA" else "CC"
    conn = get_db()
    next_num = get_next_consecutivo(conn, emisor_id, tipo, pref)
    conn.close()
    return {"next_numero": next_num, "prefijo": pref}


@app.get("/documentos/next-number", response_class=JSONResponse)
def documentos_next_number(emisor_id: int | None = None, tipo: str | None = None, prefijo: str | None = None):
    if not emisor_id or not tipo:
        return {"consecutivo": None, "prefijo": None, "display": "-"}
    pref = (prefijo or "").strip()
    if not pref:
        pref = "FV" if tipo == "FACTURA" else "CC"
    conn = get_db()
    next_num = get_next_consecutivo(conn, emisor_id, tipo, pref)
    conn.close()
    display = f"{pref}-{str(next_num).zfill(6)}"
    return {"consecutivo": next_num, "prefijo": pref, "display": display}


@app.get("/emisores", response_class=HTMLResponse)
def emisores_list(request: Request):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM emisores ORDER BY id")
    rows = cur.fetchall()
    conn.close()
    msg = request.query_params.get("msg")
    error = request.query_params.get("error")
    return templates.TemplateResponse(
        "emisores_list.html",
        {"request": request, "emisores": rows, "msg": msg, "error": error},
    )


@app.get("/emisores/nuevo", response_class=HTMLResponse)
def emisores_new_form(request: Request):
    return templates.TemplateResponse(
        "emisor_form.html",
        {"request": request, "emisor": None, "error": None},
    )


@app.post("/emisores/nuevo")
async def emisores_new(request: Request):
    form = await request.form()
    ciudad = (form.get("ciudad") or "").strip()
    departamento = (form.get("departamento") or "").strip()
    tipo_identificacion = (form.get("tipo_identificacion") or "NIT").strip()
    nombre = (form.get("nombre") or "").strip()
    identificacion = digits_only(form.get("identificacion"))
    if not nombre or not identificacion:
        return HTMLResponse(
            "Nombre e identificacion son obligatorios.",
            status_code=400,
        )
    if not ciudad:
        emisor = {
            "nombre": nombre,
            "identificacion": identificacion,
            "ciudad": ciudad,
            "departamento": departamento,
            "pais": form.get("pais") or "Colombia",
            "direccion": form.get("direccion") or "",
            "telefono": form.get("telefono") or "",
            "email": form.get("email") or "",
            "tipo_identificacion": tipo_identificacion,
        }
        return templates.TemplateResponse(
            "emisor_form.html",
            {
                "request": request,
                "emisor": emisor,
                "error": "Ciudad es obligatoria.",
            },
        )

    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        """
        INSERT INTO emisores (
            nombre, identificacion, ciudad, departamento, pais,
            direccion, telefono, email, tipo_identificacion
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            nombre,
            identificacion,
            ciudad,
            departamento,
            form.get("pais") or "Colombia",
            form.get("direccion"),
            form.get("telefono"),
            form.get("email"),
            tipo_identificacion,
        ),
    )
    conn.commit()
    conn.close()
    return RedirectResponse("/emisores", status_code=303)


@app.get("/emisores/editar/{emisor_id}", response_class=HTMLResponse)
def emisores_edit_form(request: Request, emisor_id: int):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM emisores WHERE id = ?", (emisor_id,))
    emisor = cur.fetchone()
    conn.close()
    return templates.TemplateResponse(
        "emisor_form.html",
        {"request": request, "emisor": emisor, "error": None},
    )


@app.post("/emisores/editar/{emisor_id}")
async def emisores_edit(request: Request, emisor_id: int):
    form = await request.form()
    ciudad = (form.get("ciudad") or "").strip()
    departamento = (form.get("departamento") or "").strip()
    tipo_identificacion = (form.get("tipo_identificacion") or "NIT").strip()
    if not ciudad:
        emisor = {
            "id": emisor_id,
            "nombre": form.get("nombre") or "",
            "identificacion": form.get("identificacion") or "",
            "ciudad": ciudad,
            "departamento": departamento,
            "pais": form.get("pais") or "Colombia",
            "direccion": form.get("direccion") or "",
            "telefono": form.get("telefono") or "",
            "email": form.get("email") or "",
            "tipo_identificacion": tipo_identificacion,
        }
        return templates.TemplateResponse(
            "emisor_form.html",
            {
                "request": request,
                "emisor": emisor,
                "error": "Ciudad es obligatoria.",
            },
        )

    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        """
        UPDATE emisores
        SET nombre = ?, identificacion = ?, ciudad = ?, departamento = ?, pais = ?,
            direccion = ?, telefono = ?, email = ?, tipo_identificacion = ?
        WHERE id = ?
        """,
        (
            form.get("nombre"),
            digits_only(form.get("identificacion")),
            ciudad,
            departamento,
            form.get("pais") or "Colombia",
            form.get("direccion"),
            form.get("telefono"),
            form.get("email"),
            tipo_identificacion,
            emisor_id,
        ),
    )
    conn.commit()
    conn.close()
    return RedirectResponse("/emisores", status_code=303)


@app.post("/emisores/eliminar/{emisor_id}")
def emisores_delete(emisor_id: int):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT COUNT(*) AS c FROM documentos WHERE emisor_id = ?",
        (emisor_id,),
    )
    if cur.fetchone()["c"] > 0:
        conn.close()
        msg = quote("No se puede eliminar: existen documentos asociados.")
        return RedirectResponse(f"/emisores?error={msg}", status_code=303)

    cur.execute("DELETE FROM emisores WHERE id = ?", (emisor_id,))
    conn.commit()
    conn.close()
    msg = quote("Eliminado correctamente.")
    return RedirectResponse(f"/emisores?msg={msg}", status_code=303)


@app.get("/receptores", response_class=HTMLResponse)
def receptores_list(request: Request):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM receptores ORDER BY id")
    rows = cur.fetchall()
    conn.close()
    msg = request.query_params.get("msg")
    error = request.query_params.get("error")
    return templates.TemplateResponse(
        "receptores_list.html",
        {"request": request, "receptores": rows, "msg": msg, "error": error},
    )


@app.get("/receptores/nuevo", response_class=HTMLResponse)
def receptores_new_form(request: Request):
    return templates.TemplateResponse(
        "receptor_form.html",
        {"request": request, "receptor": None, "error": None},
    )


@app.post("/receptores/nuevo")
async def receptores_new(request: Request):
    form = await request.form()
    ciudad = (form.get("ciudad") or "").strip()
    departamento = (form.get("departamento") or "").strip()
    tipo_identificacion = (form.get("tipo_identificacion") or "CC").strip()
    nombre = (form.get("nombre") or "").strip()
    identificacion = digits_only(form.get("identificacion"))
    if not nombre or not identificacion:
        return HTMLResponse(
            "Nombre e identificacion son obligatorios.",
            status_code=400,
        )
    if not ciudad:
        receptor = {
            "nombre": nombre,
            "identificacion": identificacion,
            "ciudad": ciudad,
            "departamento": departamento,
            "pais": form.get("pais") or "Colombia",
            "direccion": form.get("direccion") or "",
            "telefono": form.get("telefono") or "",
            "email": form.get("email") or "",
            "tipo_identificacion": tipo_identificacion,
        }
        return templates.TemplateResponse(
            "receptor_form.html",
            {
                "request": request,
                "receptor": receptor,
                "error": "Ciudad es obligatoria.",
            },
        )

    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        """
        INSERT INTO receptores (
            nombre, identificacion, ciudad, departamento, pais,
            direccion, telefono, email, tipo_identificacion
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            nombre,
            identificacion,
            ciudad,
            departamento,
            form.get("pais") or "Colombia",
            form.get("direccion"),
            form.get("telefono"),
            form.get("email"),
            tipo_identificacion,
        ),
    )
    conn.commit()
    conn.close()
    return RedirectResponse("/receptores", status_code=303)


@app.get("/receptores/editar/{receptor_id}", response_class=HTMLResponse)
def receptores_edit_form(request: Request, receptor_id: int):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM receptores WHERE id = ?", (receptor_id,))
    receptor = cur.fetchone()
    conn.close()
    return templates.TemplateResponse(
        "receptor_form.html",
        {"request": request, "receptor": receptor, "error": None},
    )


@app.post("/receptores/editar/{receptor_id}")
async def receptores_edit(request: Request, receptor_id: int):
    form = await request.form()
    ciudad = (form.get("ciudad") or "").strip()
    departamento = (form.get("departamento") or "").strip()
    tipo_identificacion = (form.get("tipo_identificacion") or "CC").strip()
    nombre = (form.get("nombre") or "").strip()
    identificacion = digits_only(form.get("identificacion"))
    if not nombre or not identificacion:
        return HTMLResponse(
            "Nombre e identificacion son obligatorios.",
            status_code=400,
        )
    if not ciudad:
        receptor = {
            "id": receptor_id,
            "nombre": nombre,
            "identificacion": identificacion,
            "ciudad": ciudad,
            "departamento": departamento,
            "pais": form.get("pais") or "Colombia",
            "direccion": form.get("direccion") or "",
            "telefono": form.get("telefono") or "",
            "email": form.get("email") or "",
            "tipo_identificacion": tipo_identificacion,
        }
        return templates.TemplateResponse(
            "receptor_form.html",
            {
                "request": request,
                "receptor": receptor,
                "error": "Ciudad es obligatoria.",
            },
        )

    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        """
        UPDATE receptores
        SET nombre = ?, identificacion = ?, ciudad = ?, departamento = ?, pais = ?,
            direccion = ?, telefono = ?, email = ?, tipo_identificacion = ?
        WHERE id = ?
        """,
        (
            nombre,
            identificacion,
            ciudad,
            departamento,
            form.get("pais") or "Colombia",
            form.get("direccion"),
            form.get("telefono"),
            form.get("email"),
            tipo_identificacion,
            receptor_id,
        ),
    )
    conn.commit()
    conn.close()
    return RedirectResponse("/receptores", status_code=303)


@app.post("/receptores/eliminar/{receptor_id}")
def receptores_delete(receptor_id: int):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT COUNT(*) AS c FROM documentos WHERE receptor_id = ?",
        (receptor_id,),
    )
    if cur.fetchone()["c"] > 0:
        conn.close()
        msg = quote("No se puede eliminar: existen documentos asociados.")
        return RedirectResponse(f"/receptores?error={msg}", status_code=303)

    cur.execute("DELETE FROM receptores WHERE id = ?", (receptor_id,))
    conn.commit()
    conn.close()
    msg = quote("Eliminado correctamente.")
    return RedirectResponse(f"/receptores?msg={msg}", status_code=303)


@app.get("/items", response_class=HTMLResponse)
def items_list(request: Request):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        """
        SELECT i.*, g.nombre AS grupo_nombre
        FROM items i
        LEFT JOIN grupos g ON g.id = i.grupo_id
        ORDER BY i.id
        """
    )
    rows = cur.fetchall()
    conn.close()
    msg = request.query_params.get("msg")
    error = request.query_params.get("error")
    return templates.TemplateResponse(
        "items_list.html",
        {"request": request, "items": rows, "msg": msg, "error": error},
    )


@app.get("/items/nuevo", response_class=HTMLResponse)
def items_new_form(request: Request):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM grupos ORDER BY nombre")
    grupos = cur.fetchall()
    conn.close()
    return templates.TemplateResponse(
        "item_form.html",
        {"request": request, "item": None, "grupos": grupos, "error": None},
    )


@app.post("/items/nuevo")
async def items_new(request: Request):
    form = await request.form()
    grupo_id = form.get("grupo_id")
    grupo_id_val = int(grupo_id) if grupo_id else None
    codigo = (form.get("codigo") or "").strip()
    if not codigo:
        conn = get_db()
        cur = conn.cursor()
        cur.execute("SELECT * FROM grupos ORDER BY nombre")
        grupos = cur.fetchall()
        conn.close()
        item = {
            "codigo": "",
            "descripcion": form.get("descripcion") or "",
            "costo": form.get("costo") or "",
            "precio": form.get("precio") or "",
            "iva_rate": form.get("iva_rate") or "0.00",
            "grupo_id": grupo_id_val,
        }
        return templates.TemplateResponse(
            "item_form.html",
            {
                "request": request,
                "item": item,
                "grupos": grupos,
                "error": "Codigo es obligatorio.",
            },
        )
    costo = parse_cop(form.get("costo") or 0)
    if costo < 0:
        conn = get_db()
        cur = conn.cursor()
        cur.execute("SELECT * FROM grupos ORDER BY nombre")
        grupos = cur.fetchall()
        conn.close()
        item = {
            "codigo": codigo,
            "descripcion": form.get("descripcion") or "",
            "costo": form.get("costo") or "",
            "precio": form.get("precio") or "",
            "iva_rate": form.get("iva_rate") or "0.00",
            "grupo_id": grupo_id_val,
        }
        return templates.TemplateResponse(
            "item_form.html",
            {
                "request": request,
                "item": item,
                "grupos": grupos,
                "error": "Costo debe ser un numero mayor o igual a 0.",
            },
        )
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT id FROM items WHERE lower(codigo) = lower(?)", (codigo,))
    if cur.fetchone():
        cur.execute("SELECT * FROM grupos ORDER BY nombre")
        grupos = cur.fetchall()
        conn.close()
        item = {
            "codigo": codigo,
            "descripcion": form.get("descripcion") or "",
            "costo": form.get("costo") or "",
            "precio": form.get("precio") or "",
            "iva_rate": form.get("iva_rate") or "0.00",
            "grupo_id": grupo_id_val,
        }
        return templates.TemplateResponse(
            "item_form.html",
            {
                "request": request,
                "item": item,
                "grupos": grupos,
                "error": "Codigo ya existe.",
            },
        )
    if ITEM_HAS_NOMBRE:
        cur.execute(
            """
            INSERT INTO items (nombre, codigo, descripcion, costo, precio, iva_rate, grupo_id)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                codigo,
                codigo,
                form.get("descripcion"),
                float(costo),
                float(parse_cop(form.get("precio") or 0)),
                (float(form.get("iva_rate") or 0) / 100.0),
                grupo_id_val,
            ),
        )
    else:
        cur.execute(
            """
            INSERT INTO items (codigo, descripcion, costo, precio, iva_rate, grupo_id)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                codigo,
                form.get("descripcion"),
                float(costo),
                float(parse_cop(form.get("precio") or 0)),
                (float(form.get("iva_rate") or 0) / 100.0),
                grupo_id_val,
            ),
        )
    conn.commit()
    conn.close()
    return RedirectResponse("/items", status_code=303)


@app.get("/items/editar/{item_id}", response_class=HTMLResponse)
def items_edit_form(request: Request, item_id: int):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM items WHERE id = ?", (item_id,))
    item = cur.fetchone()
    cur.execute("SELECT * FROM grupos ORDER BY nombre")
    grupos = cur.fetchall()
    conn.close()
    return templates.TemplateResponse(
        "item_form.html",
        {"request": request, "item": item, "grupos": grupos, "error": None},
    )


@app.post("/items/editar/{item_id}")
async def items_edit(request: Request, item_id: int):
    form = await request.form()
    grupo_id = form.get("grupo_id")
    grupo_id_val = int(grupo_id) if grupo_id else None
    codigo = (form.get("codigo") or "").strip()
    if not codigo:
        conn = get_db()
        cur = conn.cursor()
        cur.execute("SELECT * FROM grupos ORDER BY nombre")
        grupos = cur.fetchall()
        conn.close()
        item = {
            "id": item_id,
            "codigo": "",
            "descripcion": form.get("descripcion") or "",
            "costo": form.get("costo") or "",
            "precio": form.get("precio") or "",
            "iva_rate": form.get("iva_rate") or "0.00",
            "grupo_id": grupo_id_val,
        }
        return templates.TemplateResponse(
            "item_form.html",
            {
                "request": request,
                "item": item,
                "grupos": grupos,
                "error": "Codigo es obligatorio.",
            },
        )
    costo = parse_cop(form.get("costo") or 0)
    if costo < 0:
        conn = get_db()
        cur = conn.cursor()
        cur.execute("SELECT * FROM grupos ORDER BY nombre")
        grupos = cur.fetchall()
        conn.close()
        item = {
            "id": item_id,
            "codigo": codigo,
            "descripcion": form.get("descripcion") or "",
            "costo": form.get("costo") or "",
            "precio": form.get("precio") or "",
            "iva_rate": form.get("iva_rate") or "0.00",
            "grupo_id": grupo_id_val,
        }
        return templates.TemplateResponse(
            "item_form.html",
            {
                "request": request,
                "item": item,
                "grupos": grupos,
                "error": "Costo debe ser un numero mayor o igual a 0.",
            },
        )
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT id FROM items WHERE lower(codigo) = lower(?) AND id != ?",
        (codigo, item_id),
    )
    if cur.fetchone():
        cur.execute("SELECT * FROM grupos ORDER BY nombre")
        grupos = cur.fetchall()
        conn.close()
        item = {
            "id": item_id,
            "codigo": codigo,
            "descripcion": form.get("descripcion") or "",
            "costo": form.get("costo") or "",
            "precio": form.get("precio") or "",
            "iva_rate": form.get("iva_rate") or "0.00",
            "grupo_id": grupo_id_val,
        }
        return templates.TemplateResponse(
            "item_form.html",
            {
                "request": request,
                "item": item,
                "grupos": grupos,
                "error": "Codigo ya existe.",
            },
        )
    if ITEM_HAS_NOMBRE:
        cur.execute(
            """
            UPDATE items
            SET nombre = ?, codigo = ?, descripcion = ?, costo = ?, precio = ?, iva_rate = ?, grupo_id = ?
            WHERE id = ?
            """,
            (
                codigo,
                codigo,
                form.get("descripcion"),
                float(costo),
                float(parse_cop(form.get("precio") or 0)),
                (float(form.get("iva_rate") or 0) / 100.0),
                grupo_id_val,
                item_id,
            ),
        )
    else:
        cur.execute(
            """
            UPDATE items
            SET codigo = ?, descripcion = ?, costo = ?, precio = ?, iva_rate = ?, grupo_id = ?
            WHERE id = ?
            """,
            (
                codigo,
                form.get("descripcion"),
                float(costo),
                float(parse_cop(form.get("precio") or 0)),
                (float(form.get("iva_rate") or 0) / 100.0),
                grupo_id_val,
                item_id,
            ),
        )
    conn.commit()
    conn.close()
    return RedirectResponse("/items", status_code=303)


@app.post("/items/eliminar/{item_id}")
def items_delete(item_id: int):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT COUNT(*) AS c FROM documento_detalle WHERE item_id = ?",
        (item_id,),
    )
    if cur.fetchone()["c"] > 0:
        conn.close()
        msg = quote("No se puede eliminar: existen documentos asociados.")
        return RedirectResponse(f"/items?error={msg}", status_code=303)

    cur.execute("DELETE FROM items WHERE id = ?", (item_id,))
    conn.commit()
    conn.close()
    msg = quote("Eliminado correctamente.")
    return RedirectResponse(f"/items?msg={msg}", status_code=303)


@app.get("/grupos", response_class=HTMLResponse)
def grupos_list(request: Request):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM grupos ORDER BY nombre")
    rows = cur.fetchall()
    conn.close()
    msg = request.query_params.get("msg")
    error = request.query_params.get("error")
    return templates.TemplateResponse(
        "grupos_list.html",
        {"request": request, "grupos": rows, "msg": msg, "error": error},
    )


@app.get("/grupos/nuevo", response_class=HTMLResponse)
def grupos_new_form(request: Request):
    return templates.TemplateResponse(
        "grupos_form.html",
        {"request": request, "grupo": None, "error": None},
    )


@app.post("/grupos/nuevo")
async def grupos_new(request: Request):
    form = await request.form()
    nombre = (form.get("nombre") or "").strip()
    if not nombre:
        return templates.TemplateResponse(
            "grupos_form.html",
            {"request": request, "grupo": {"nombre": ""}, "error": "Nombre es obligatorio."},
        )

    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT id FROM grupos WHERE lower(nombre) = lower(?)", (nombre,))
    row = cur.fetchone()
    if row:
        conn.close()
        msg = quote("Grupo ya existe. Se usa el existente.")
        return RedirectResponse(f"/grupos?msg={msg}", status_code=303)

    cur.execute("INSERT INTO grupos (nombre) VALUES (?)", (nombre,))
    conn.commit()
    conn.close()
    msg = quote("Grupo creado.")
    return RedirectResponse(f"/grupos?msg={msg}", status_code=303)


@app.get("/grupos/editar/{grupo_id}", response_class=HTMLResponse)
def grupos_edit_form(request: Request, grupo_id: int):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM grupos WHERE id = ?", (grupo_id,))
    grupo = cur.fetchone()
    conn.close()
    return templates.TemplateResponse(
        "grupos_form.html",
        {"request": request, "grupo": grupo, "error": None},
    )


@app.post("/grupos/editar/{grupo_id}")
async def grupos_edit(request: Request, grupo_id: int):
    form = await request.form()
    nombre = (form.get("nombre") or "").strip()
    if not nombre:
        return templates.TemplateResponse(
            "grupos_form.html",
            {"request": request, "grupo": {"id": grupo_id, "nombre": ""}, "error": "Nombre es obligatorio."},
        )

    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "SELECT id FROM grupos WHERE lower(nombre) = lower(?) AND id != ?",
        (nombre, grupo_id),
    )
    if cur.fetchone():
        conn.close()
        return templates.TemplateResponse(
            "grupos_form.html",
            {
                "request": request,
                "grupo": {"id": grupo_id, "nombre": nombre},
                "error": "Ya existe un grupo con ese nombre.",
            },
        )

    cur.execute("UPDATE grupos SET nombre = ? WHERE id = ?", (nombre, grupo_id))
    conn.commit()
    conn.close()
    msg = quote("Grupo actualizado.")
    return RedirectResponse(f"/grupos?msg={msg}", status_code=303)


@app.post("/grupos/eliminar/{grupo_id}")
def grupos_delete(grupo_id: int):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) AS c FROM items WHERE grupo_id = ?", (grupo_id,))
    if cur.fetchone()["c"] > 0:
        conn.close()
        msg = quote("No se puede eliminar: hay items asociados.")
        return RedirectResponse(f"/grupos?error={msg}", status_code=303)

    cur.execute("DELETE FROM grupos WHERE id = ?", (grupo_id,))
    conn.commit()
    conn.close()
    msg = quote("Grupo eliminado.")
    return RedirectResponse(f"/grupos?msg={msg}", status_code=303)


@app.get("/documentos/nuevo", response_class=HTMLResponse)
def documentos_new_form(request: Request):
    conn = get_db()
    cur = conn.cursor()
    u = request.session.get("user") or {}
    if u.get("role") == "cliente_receptor":
        cur.execute("SELECT * FROM emisores WHERE id = ? ORDER BY nombre", (u.get("locked_emisor_id"),))
        emisores = cur.fetchall()
        cur.execute("SELECT * FROM receptores WHERE id = ? ORDER BY nombre", (u.get("locked_receptor_id"),))
        receptores = cur.fetchall()
    else:
        u = request.session.get("user") or {}
        if u.get("role") == "cliente_receptor":
            cur.execute("SELECT * FROM emisores WHERE id = ? ORDER BY nombre", (u.get("locked_emisor_id"),))
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores WHERE id = ? ORDER BY nombre", (u.get("locked_receptor_id"),))
            receptores = cur.fetchall()
        else:
            cur.execute("SELECT * FROM emisores ORDER BY nombre")
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores ORDER BY nombre")
            receptores = cur.fetchall()
    cur.execute("SELECT * FROM items ORDER BY descripcion, codigo")
    items = cur.fetchall()
    conn.close()
    today = datetime.date.today().isoformat()
    msg = request.query_params.get("msg")
    response = templates.TemplateResponse(
        "documento_form.html",
        {
            "request": request,
            "emisores": emisores,
            "receptores": receptores,
            "items": items,
            "error": None,
            "msg": msg,
            "fecha": today,
            "draft": None,
            "lines": None,
            "emit_action": None,
            "borrador_action": None,
        },
    )
    response.headers["Cache-Control"] = "no-store"
    return response


@app.post("/documentos/nuevo")
async def documentos_new(request: Request):
    form = await request.form()
    emisor_id = int(form.get("emisor_id"))
    receptor_id = int(form.get("receptor_id"))
    tipo = form.get("tipo")
    prefijo = (form.get("prefijo") or "").strip()
    if not prefijo:
        prefijo = "FV" if tipo == "FACTURA" else "CC"
    fecha = (form.get("fecha") or "").strip()
    if not fecha:
        fecha = datetime.date.today().isoformat()
    numero_sugerido = (form.get("numero_sugerido") or "").strip()
    observaciones = (form.get("observaciones") or "").strip()
    include_tax_note_383 = (
        parse_checkbox_value(form.get("include_tax_note_383")) if tipo == "CUENTA_COBRO" else 0
    )

    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM items")
    items_rows = {row["id"]: row for row in cur.fetchall()}
    doc_lines = build_lines_from_form(form, items_rows)
    if not doc_lines:
        u = request.session.get("user") or {}
        if u.get("role") == "cliente_receptor":
            cur.execute("SELECT * FROM emisores WHERE id = ? ORDER BY nombre", (u.get("locked_emisor_id"),))
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores WHERE id = ? ORDER BY nombre", (u.get("locked_receptor_id"),))
            receptores = cur.fetchall()
        else:
            cur.execute("SELECT * FROM emisores ORDER BY nombre")
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores ORDER BY nombre")
            receptores = cur.fetchall()
        cur.execute("SELECT * FROM items ORDER BY descripcion, codigo")
        items = cur.fetchall()
        conn.close()
        return templates.TemplateResponse(
            "documento_form.html",
            {
                "request": request,
                "emisores": emisores,
                "receptores": receptores,
                "items": items,
                "error": "Debe agregar al menos una linea con cantidad.",
                "msg": None,
                "fecha": fecha,
                "draft": None,
                "lines": None,
                "emit_action": None,
                "borrador_action": None,
            },
        )

    if not numero_sugerido.isdigit():
        u = request.session.get("user") or {}
        if u.get("role") == "cliente_receptor":
            cur.execute("SELECT * FROM emisores WHERE id = ? ORDER BY nombre", (u.get("locked_emisor_id"),))
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores WHERE id = ? ORDER BY nombre", (u.get("locked_receptor_id"),))
            receptores = cur.fetchall()
        else:
            cur.execute("SELECT * FROM emisores ORDER BY nombre")
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores ORDER BY nombre")
            receptores = cur.fetchall()
        cur.execute("SELECT * FROM items ORDER BY descripcion, codigo")
        items = cur.fetchall()
        conn.close()
        return templates.TemplateResponse(
            "documento_form.html",
            {
                "request": request,
                "emisores": emisores,
                "receptores": receptores,
                "items": items,
                "error": "Debe seleccionar Emisor y Tipo para asignar el consecutivo.",
                "msg": None,
                "fecha": fecha,
                "draft": None,
                "lines": None,
                "emit_action": None,
                "borrador_action": None,
            },
        )

    if numero_sugerido.isdigit():
        next_num = get_next_consecutivo(conn, emisor_id, tipo, prefijo)
        if int(numero_sugerido) != next_num:
            cur.execute("SELECT * FROM emisores ORDER BY nombre")
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores ORDER BY nombre")
            receptores = cur.fetchall()
            cur.execute("SELECT * FROM items ORDER BY descripcion, codigo")
            items = cur.fetchall()
            conn.close()
            return templates.TemplateResponse(
                "documento_form.html",
                {
                    "request": request,
                    "emisores": emisores,
                    "receptores": receptores,
                    "items": items,
                    "error": "El consecutivo cambio. Actualiza el formulario y vuelve a intentar.",
                    "msg": None,
                    "fecha": fecha,
                    "draft": None,
                    "lines": None,
                    "emit_action": None,
                    "borrador_action": None,
                },
            )

    doc_id, _numero = create_document(
        conn,
        emisor_id,
        receptor_id,
        tipo,
        prefijo,
        doc_lines,
        fecha,
        observaciones,
        include_tax_note_383,
    )
    conn.close()
    return RedirectResponse(f"/documentos/{doc_id}", status_code=303)


@app.post("/documentos/borrador")
async def documentos_borrador(request: Request):
    form = await request.form()
    emisor_id = int(form.get("emisor_id"))
    receptor_id = int(form.get("receptor_id"))
    tipo = form.get("tipo")
    prefijo = (form.get("prefijo") or "").strip()
    if not prefijo:
        prefijo = "FV" if tipo == "FACTURA" else "CC"
    fecha = (form.get("fecha") or "").strip()
    if not fecha:
        fecha = datetime.date.today().isoformat()
    numero_sugerido = (form.get("numero_sugerido") or "").strip()
    observaciones = (form.get("observaciones") or "").strip()
    include_tax_note_383 = (
        parse_checkbox_value(form.get("include_tax_note_383")) if tipo == "CUENTA_COBRO" else 0
    )

    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT * FROM items")
    items_rows = {row["id"]: row for row in cur.fetchall()}
    doc_lines = build_lines_from_form(form, items_rows)
    if not doc_lines:
        conn.close()
        msg = quote("No hay lineas para guardar borrador.")
        return RedirectResponse(f"/documentos/nuevo?msg={msg}", status_code=303)

    if not numero_sugerido.isdigit():
        u = request.session.get("user") or {}
        if u.get("role") == "cliente_receptor":
            cur.execute("SELECT * FROM emisores WHERE id = ? ORDER BY nombre", (u.get("locked_emisor_id"),))
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores WHERE id = ? ORDER BY nombre", (u.get("locked_receptor_id"),))
            receptores = cur.fetchall()
        else:
            cur.execute("SELECT * FROM emisores ORDER BY nombre")
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores ORDER BY nombre")
            receptores = cur.fetchall()
        cur.execute("SELECT * FROM items ORDER BY descripcion, codigo")
        items = cur.fetchall()
        conn.close()
        return templates.TemplateResponse(
            "documento_form.html",
            {
                "request": request,
                "emisores": emisores,
                "receptores": receptores,
                "items": items,
                "error": "Debe seleccionar Emisor y Tipo para asignar el consecutivo.",
                "msg": None,
                "fecha": fecha,
                "draft": None,
                "lines": None,
                "emit_action": None,
                "borrador_action": None,
            },
        )

    next_num = get_next_consecutivo(conn, emisor_id, tipo, prefijo)
    if int(numero_sugerido) != next_num:
        u = request.session.get("user") or {}
        if u.get("role") == "cliente_receptor":
            cur.execute("SELECT * FROM emisores WHERE id = ? ORDER BY nombre", (u.get("locked_emisor_id"),))
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores WHERE id = ? ORDER BY nombre", (u.get("locked_receptor_id"),))
            receptores = cur.fetchall()
        else:
            cur.execute("SELECT * FROM emisores ORDER BY nombre")
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores ORDER BY nombre")
            receptores = cur.fetchall()
        cur.execute("SELECT * FROM items ORDER BY descripcion, codigo")
        items = cur.fetchall()
        conn.close()
        return templates.TemplateResponse(
            "documento_form.html",
            {
                "request": request,
                "emisores": emisores,
                "receptores": receptores,
                "items": items,
                "error": "El consecutivo cambio. Actualiza el formulario y vuelve a intentar.",
                "msg": None,
                "fecha": fecha,
                "draft": None,
                "lines": None,
                "emit_action": None,
                "borrador_action": None,
            },
        )

    numero = get_consecutivo(conn, emisor_id, tipo, prefijo)
    borrador_id = create_draft_document(
        conn,
        emisor_id,
        receptor_id,
        tipo,
        prefijo,
        doc_lines,
        fecha,
        observaciones,
        numero,
        include_tax_note_383,
    )
    conn.close()
    return RedirectResponse(f"/documentos/borrador/{borrador_id}", status_code=303)


@app.get("/documentos/borrador/{doc_id}", response_class=HTMLResponse)
def documentos_borrador_view(request: Request, doc_id: int):
    conn = get_db()
    doc, lines = fetch_document(conn, doc_id)
    conn.close()
    if not doc or doc["estado"] != "BORRADOR":
        return HTMLResponse("Borrador no encontrado.", status_code=404)
    return templates.TemplateResponse(
        "documento_borrador_view.html",
        {"request": request, "doc": doc, "lines": lines},
    )


@app.get("/documentos/borrador/{doc_id}/editar", response_class=HTMLResponse)
def documentos_borrador_edit_form(request: Request, doc_id: int):
    return RedirectResponse(f"/documentos/editar/{doc_id}", status_code=303)


@app.post("/documentos/borrador/{doc_id}/editar")
async def documentos_borrador_edit(doc_id: int, request: Request):
    return RedirectResponse(f"/documentos/editar/{doc_id}", status_code=307)


@app.get("/documentos/editar/{doc_id}", response_class=HTMLResponse)
def documentos_edit_form(request: Request, doc_id: int):
    conn = get_db()
    doc, lines = fetch_document(conn, doc_id)
    if not doc:
        conn.close()
        return HTMLResponse("Documento no encontrado.", status_code=404)
    cur = conn.cursor()
    u = request.session.get("user") or {}
    if u.get("role") == "cliente_receptor":
        cur.execute("SELECT * FROM emisores WHERE id = ? ORDER BY nombre", (u.get("locked_emisor_id"),))
        emisores = cur.fetchall()
        cur.execute("SELECT * FROM receptores WHERE id = ? ORDER BY nombre", (u.get("locked_receptor_id"),))
        receptores = cur.fetchall()
    else:
        u = request.session.get("user") or {}
        if u.get("role") == "cliente_receptor":
            cur.execute("SELECT * FROM emisores WHERE id = ? ORDER BY nombre", (u.get("locked_emisor_id"),))
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores WHERE id = ? ORDER BY nombre", (u.get("locked_receptor_id"),))
            receptores = cur.fetchall()
        else:
            cur.execute("SELECT * FROM emisores ORDER BY nombre")
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores ORDER BY nombre")
            receptores = cur.fetchall()
    cur.execute("SELECT * FROM items ORDER BY descripcion, codigo")
    items = cur.fetchall()
    conn.close()

    filled_lines = []
    for line in lines:
        filled_lines.append(
            {
                "item_id": line["item_id"],
                "cantidad": line["cantidad"],
                "precio_unitario": line["precio_unitario"],
                "iva_percent": float(line["iva_rate"]) * 100,
            }
        )
    if not filled_lines:
        filled_lines.append({})

    msg = request.query_params.get("msg")
    response = templates.TemplateResponse(
        "documento_form.html",
        {
            "request": request,
            "emisores": emisores,
            "receptores": receptores,
            "items": items,
            "error": None,
            "msg": msg,
            "fecha": doc["fecha"],
            "draft": doc,
            "lines": filled_lines,
            "emit_action": f"/documentos/borrador/{doc_id}/emitir" if doc["estado"] == "BORRADOR" else None,
            "borrador_action": f"/documentos/editar/{doc_id}",
            "edit_mode": True,
        },
    )
    response.headers["Cache-Control"] = "no-store"
    return response


@app.post("/documentos/editar/{doc_id}")
async def documentos_edit(doc_id: int, request: Request):
    form = await request.form()
    conn = get_db()
    doc, _lines = fetch_document(conn, doc_id)
    if not doc:
        conn.close()
        return HTMLResponse("Documento no encontrado.", status_code=404)

    estado = doc["estado"]
    if estado == "EMITIDO":
        emisor_id = doc["emisor_id"]
        receptor_id = doc["receptor_id"]
        tipo = doc["tipo"]
        prefijo = doc["prefijo"]
        numero = doc["numero"]
    else:
        emisor_id = int(form.get("emisor_id"))
        receptor_id = int(form.get("receptor_id"))
        tipo = form.get("tipo")
        prefijo = (form.get("prefijo") or "").strip()
        if not prefijo:
            prefijo = "FV" if tipo == "FACTURA" else "CC"
        numero_sugerido = (form.get("numero_sugerido") or "").strip()
        if (
            doc["numero"] is not None
            and numero_sugerido.isdigit()
            and int(numero_sugerido) == doc["numero"]
            and emisor_id == doc["emisor_id"]
            and tipo == doc["tipo"]
            and prefijo == doc["prefijo"]
        ):
            numero = doc["numero"]
        else:
            if not numero_sugerido.isdigit():
                conn.close()
                msg = quote("Debe seleccionar Emisor y Tipo para asignar el consecutivo.")
                return RedirectResponse(f"/documentos/editar/{doc_id}?msg={msg}", status_code=303)
            next_num = get_next_consecutivo(conn, emisor_id, tipo, prefijo)
            if int(numero_sugerido) != next_num:
                conn.close()
                msg = quote("El consecutivo cambio. Actualiza el formulario y vuelve a intentar.")
                return RedirectResponse(f"/documentos/editar/{doc_id}?msg={msg}", status_code=303)
            numero = get_consecutivo(conn, emisor_id, tipo, prefijo)

    fecha = (form.get("fecha") or "").strip()
    if not fecha:
        fecha = datetime.date.today().isoformat()
    observaciones = (form.get("observaciones") or "").strip()
    include_tax_note_383 = (
        parse_checkbox_value(form.get("include_tax_note_383")) if tipo == "CUENTA_COBRO" else 0
    )

    cur = conn.cursor()
    cur.execute("SELECT * FROM items")
    items_rows = {row["id"]: row for row in cur.fetchall()}
    doc_lines = build_lines_from_form(form, items_rows)
    if not doc_lines:
        conn.close()
        msg = quote("Debe agregar al menos una linea con cantidad.")
        return RedirectResponse(f"/documentos/editar/{doc_id}?msg={msg}", status_code=303)

    subtotal, iva, total = calculate_totals(doc_lines)
    created_at = datetime.datetime.now().isoformat(timespec="seconds")

    cur.execute(
        """
        UPDATE documentos
        SET emisor_id = ?, receptor_id = ?, tipo = ?, prefijo = ?, numero = ?,
            fecha = ?, subtotal = ?, iva = ?, total = ?, observaciones = ?, include_tax_note_383 = ?, created_at = ?, estado = ?
        WHERE id = ?
        """,
        (
            emisor_id,
            receptor_id,
            tipo,
            prefijo,
            numero,
            fecha,
            to_sqlite_number(subtotal),
            to_sqlite_number(iva),
            to_sqlite_number(total),
            observaciones,
            include_tax_note_383,
            created_at,
            estado,
            doc_id,
        ),
    )
    cur.execute("DELETE FROM documento_detalle WHERE documento_id = ?", (doc_id,))
    for line in doc_lines:
        cur.execute(
            """
            INSERT INTO documento_detalle (
                documento_id, item_id, descripcion, cantidad, precio_unitario,
                iva_rate, subtotal, iva, total
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                doc_id,
                line.get("item_id"),
                line["descripcion"],
                to_sqlite_number(line["cantidad"]),
                to_sqlite_number(line["precio_unitario"]),
                to_sqlite_number(line["iva_rate"]),
                to_sqlite_number(line["subtotal"]),
                to_sqlite_number(line["iva"]),
                to_sqlite_number(line["total"]),
            ),
        )
    conn.commit()
    conn.close()
    return RedirectResponse(f"/documentos/{doc_id}" if estado == "EMITIDO" else f"/documentos/borrador/{doc_id}", status_code=303)


@app.post("/documentos/borrador/{doc_id}/emitir")
def documentos_borrador_emitir(doc_id: int):
    conn = get_db()
    doc, lines = fetch_document(conn, doc_id)
    if not doc or doc["estado"] != "BORRADOR":
        conn.close()
        return HTMLResponse("Borrador no encontrado.", status_code=404)

    doc_lines = []
    for line in lines:
        doc_lines.append(
            {
                "id": line["id"],
                "item_id": line["item_id"],
                "descripcion": line["descripcion"],
                "cantidad": float(line["cantidad"]),
                "precio_unitario": float(line["precio_unitario"]),
                "iva_rate": float(line["iva_rate"]),
            }
        )
    subtotal, iva, total = calculate_totals(doc_lines)
    if doc["numero"] is not None:
        numero = doc["numero"]
    else:
        numero = get_consecutivo(conn, doc["emisor_id"], doc["tipo"], doc["prefijo"])
    created_at = datetime.datetime.now().isoformat(timespec="seconds")

    cur = conn.cursor()
    cur.execute(
        """
        UPDATE documentos
        SET numero = ?, subtotal = ?, iva = ?, total = ?, created_at = ?, estado = 'EMITIDO'
        WHERE id = ?
        """,
        (
            numero,
            to_sqlite_number(subtotal),
            to_sqlite_number(iva),
            to_sqlite_number(total),
            created_at,
            doc_id,
        ),
    )
    for line in doc_lines:
        cur.execute(
            """
            UPDATE documento_detalle
            SET subtotal = ?, iva = ?, total = ?
            WHERE id = ?
            """,
            (
                to_sqlite_number(line["subtotal"]),
                to_sqlite_number(line["iva"]),
                to_sqlite_number(line["total"]),
                line["id"],
            ),
        )
    conn.commit()
    conn.close()
    return RedirectResponse(f"/documentos/{doc_id}", status_code=303)


@app.get("/documentos/consultar", response_class=HTMLResponse)
def documentos_consulta(request: Request):
    conn = get_db()
    cur = conn.cursor()
    u = request.session.get("user") or {}
    if u.get("role") == "cliente_receptor":
        cur.execute("SELECT * FROM emisores WHERE id = ? ORDER BY nombre", (u.get("locked_emisor_id"),))
        emisores = cur.fetchall()
        cur.execute("SELECT * FROM receptores WHERE id = ? ORDER BY nombre", (u.get("locked_receptor_id"),))
        receptores = cur.fetchall()
    else:
        u = request.session.get("user") or {}
        if u.get("role") == "cliente_receptor":
            cur.execute("SELECT * FROM emisores WHERE id = ? ORDER BY nombre", (u.get("locked_emisor_id"),))
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores WHERE id = ? ORDER BY nombre", (u.get("locked_receptor_id"),))
            receptores = cur.fetchall()
        else:
            cur.execute("SELECT * FROM emisores ORDER BY nombre")
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores ORDER BY nombre")
            receptores = cur.fetchall()

    where, params, filters = parse_document_filters(
        request.query_params, default_estado="EMITIDO"
    )
    docs = fetch_documents(conn, where, params, limit=None if where else 50)
    total_consulta = fetch_documents_sum(conn, where, params)
    conn.close()

    msg = None
    if not where:
        msg = "Mostrando los ultimos 50 documentos."
    error = None
    if where and not docs:
        error = "No se encontraron documentos con esos filtros."

    return templates.TemplateResponse(
        "documento_consulta.html",
        {
            "request": request,
            "emisores": emisores,
            "receptores": receptores,
            "docs": docs,
            "total_consulta": total_consulta,
            "error": error,
            "msg": msg,
            "filters": filters,
        },
    )


@app.post("/documentos/eliminar/{doc_id}")
def documentos_eliminar(doc_id: int):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT id FROM documentos WHERE id = ?", (doc_id,))
    row = cur.fetchone()
    if not row:
        conn.close()
        msg = quote("Documento no encontrado.")
        return RedirectResponse(f"/documentos/consultar?error={msg}", status_code=303)

    cur.execute("DELETE FROM documento_detalle WHERE documento_id = ?", (doc_id,))
    cur.execute("DELETE FROM documentos WHERE id = ?", (doc_id,))
    conn.commit()
    conn.close()
    msg = quote("Documento eliminado.")
    return RedirectResponse(f"/documentos/consultar?msg={msg}", status_code=303)


@app.get("/documentos/{doc_id}/pdf")
def documentos_pdf(doc_id: int, inline: bool = False):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        """
        SELECT d.*, e.nombre AS emisor_nombre, r.nombre AS receptor_nombre,
               e.identificacion AS emisor_identificacion, r.identificacion AS receptor_identificacion,
               e.tipo_identificacion AS emisor_tipo_identificacion,
               r.tipo_identificacion AS receptor_tipo_identificacion,
               e.ciudad AS emisor_ciudad, e.departamento AS emisor_departamento,
               r.ciudad AS receptor_ciudad, r.departamento AS receptor_departamento,
               e.direccion AS emisor_direccion, e.email AS emisor_email, e.telefono AS emisor_telefono,
               r.direccion AS receptor_direccion, r.email AS receptor_email, r.telefono AS receptor_telefono
        FROM documentos d
        JOIN emisores e ON e.id = d.emisor_id
        JOIN receptores r ON r.id = d.receptor_id
        WHERE d.id = ?
        """,
        (doc_id,),
    )
    doc = cur.fetchone()
    if not doc:
        conn.close()
        return HTMLResponse("Documento no encontrado.", status_code=404)
    if doc["estado"] == "BORRADOR":
        conn.close()
        return HTMLResponse("El documento esta en borrador.", status_code=400)
    cur.execute(
        """
        SELECT dd.*, i.descripcion AS item_descripcion, i.codigo AS item_codigo
        FROM documento_detalle dd
        LEFT JOIN items i ON i.id = dd.item_id
        WHERE dd.documento_id = ?
        ORDER BY dd.id
        """,
        (doc_id,),
    )
    lines = cur.fetchall()
    conn.close()
    pdf_buffer = build_document_pdf(doc, lines)
    tipo = "FACTURA_DE_VENTA" if doc["tipo"] == "FACTURA" else "CUENTA_COBRO"
    filename = f"{tipo}_{doc['prefijo']}-{doc['numero']:06d}.pdf"
    disposition = "inline" if inline else "attachment"
    headers = {"Content-Disposition": f'{disposition}; filename="{filename}"'}
    return StreamingResponse(pdf_buffer, media_type="application/pdf", headers=headers)


@app.get("/documentos/{doc_id}/pdf/preview", response_class=HTMLResponse)
def documentos_pdf_preview(request: Request, doc_id: int):
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT id, tipo, prefijo, numero, estado FROM documentos WHERE id = ?", (doc_id,))
    doc = cur.fetchone()
    conn.close()
    if not doc:
        return HTMLResponse("Documento no encontrado.", status_code=404)
    if doc["estado"] == "BORRADOR":
        return HTMLResponse("El documento esta en borrador.", status_code=400)
    return templates.TemplateResponse(
        "documento_pdf_preview.html",
        {"request": request, "doc": doc},
    )


@app.get("/documentos/{doc_id}", response_class=HTMLResponse)
def documentos_view(request: Request, doc_id: int):
    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        """
        SELECT d.*, e.nombre AS emisor_nombre, r.nombre AS receptor_nombre,
               e.identificacion AS emisor_identificacion, r.identificacion AS receptor_identificacion,
               e.tipo_identificacion AS emisor_tipo_identificacion,
               r.tipo_identificacion AS receptor_tipo_identificacion,
               e.ciudad AS emisor_ciudad, e.departamento AS emisor_departamento,
               r.ciudad AS receptor_ciudad, r.departamento AS receptor_departamento,
               e.direccion AS emisor_direccion, e.email AS emisor_email, e.telefono AS emisor_telefono,
               r.direccion AS receptor_direccion, r.email AS receptor_email, r.telefono AS receptor_telefono
        FROM documentos d
        JOIN emisores e ON e.id = d.emisor_id
        JOIN receptores r ON r.id = d.receptor_id
        WHERE d.id = ?
        """,
        (doc_id,),
    )
    doc = cur.fetchone()
    if not doc:
        conn.close()
        return HTMLResponse("Documento no encontrado.", status_code=404)
    cur.execute(
        """
        SELECT dd.*, i.descripcion AS item_descripcion, i.codigo AS item_codigo
        FROM documento_detalle dd
        LEFT JOIN items i ON i.id = dd.item_id
        WHERE dd.documento_id = ?
        ORDER BY dd.id
        """,
        (doc_id,),
    )
    lines = cur.fetchall()
    conn.close()
    template_name = "documento_view.html"
    if doc["estado"] == "BORRADOR":
        template_name = "documento_borrador_view.html"
    return templates.TemplateResponse(
        template_name,
        {"request": request, "doc": doc, "lines": lines},
    )


def build_informe_documentos_pdf(lines, header, totals, count_docs, mode):
    buffer = BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("Title", parent=styles["Heading1"], alignment=1, fontSize=16, leading=18)
    normal = ParagraphStyle("Normal", parent=styles["Normal"], fontSize=10, leading=14)
    label = ParagraphStyle("Label", parent=styles["Normal"], fontSize=9, textColor=colors.grey)
    desc_style = ParagraphStyle("Desc", parent=styles["Normal"], fontSize=8, leading=10)

    story = []
    story.append(Paragraph("Informe de documentos", title))
    story.append(Spacer(1, 6))
    story.append(Paragraph(f"<b>Emisor:</b> {header['emisor']}", normal))
    story.append(Paragraph(f"<b>Cliente:</b> {header['cliente']}", normal))
    story.append(Paragraph(f"<b>Rango de fechas:</b> {header['fechas']}", normal))
    story.append(Paragraph(f"<b>Tipo:</b> {header['tipo']}", normal))
    story.append(Spacer(1, 10))

    if mode == "resumen":
        table_data = [
            ["Descripcion", "Cantidad", "Base", "IVA", "Total"]
        ]
        for line in lines:
            table_data.append(
                [
                    Paragraph(str(line["descripcion"] or ""), desc_style),
                    f"{line['cantidad']:.2f}",
                    currency_cop(line["base"] or 0),
                    currency_cop(line["iva"] or 0),
                    currency_cop(line["total"] or 0),
                ]
            )
        table_data.append(
            [
                Paragraph("TOTAL", desc_style),
                f"{totals['cantidad']:.2f}",
                currency_cop(totals["base"]),
                currency_cop(totals["iva"]),
                currency_cop(totals["total"]),
            ]
        )
        col_widths = [7.0 * cm, 2.0 * cm, 3.0 * cm, 3.0 * cm, 3.0 * cm]
        align_start = 1
    else:
        table_data = [
            ["Fecha", "Tipo", "Numero", "Descripcion", "Cantidad", "Valor unitario", "IVA %", "IVA $", "Total linea"]
        ]
        for line in lines:
            numero = "-"
            if line["numero"] is not None:
                numero = f"{line['prefijo']}-{int(line['numero']):06d}"
            table_data.append(
                [
                    line["fecha"],
                    Paragraph(format_tipo_display(line["tipo"]), desc_style),
                    numero,
                    Paragraph(str(line["descripcion"] or ""), desc_style),
                    f"{line['cantidad']:.2f}",
                    currency_cop(line["precio_unitario"] or 0),
                    format_percent(line["iva_rate"] or 0),
                    currency_cop(line["iva"] or 0),
                    currency_cop(line["total"] or 0),
                ]
            )
        table_data.append(
            [
                "",
                "",
                "",
                Paragraph("TOTAL", desc_style),
                f"{totals['cantidad']:.2f}",
                "",
                "",
                currency_cop(totals["iva"]),
                currency_cop(totals["total"]),
            ]
        )
        col_widths = [
            1.6 * cm,  # Fecha
            1.5 * cm,  # Tipo
            2.2 * cm,  # Numero
            5.0 * cm,  # Descripcion
            1.2 * cm,  # Cantidad
            2.0 * cm,  # Valor unitario
            1.0 * cm,  # IVA %
            1.6 * cm,  # IVA $
            1.5 * cm,  # Total linea
        ]
        align_start = 4

    table = Table(table_data, repeatRows=1, colWidths=col_widths)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f4f8")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.lightgrey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (3, -1), "LEFT"),
                ("ALIGN", (align_start, 1), (-1, -1), "RIGHT"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"<b>Cantidad de documentos:</b> {count_docs}", normal))

    pdf.build(story)
    buffer.seek(0)
    return buffer


def build_informe_documentos_excel(lines, header, totals, count_docs, mode):
    wb = Workbook()
    ws = wb.active
    ws.title = "Documentos"

    ws["A1"] = "Informe de documentos"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generado: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A3"] = f"Emisor: {header['emisor']}"
    ws["A4"] = f"Cliente: {header['cliente']}"
    ws["A5"] = f"Rango de fechas: {header['fechas']}"
    ws["A6"] = f"Tipo: {header['tipo']}"

    headers = ["Fecha", "Tipo", "Numero", "Descripcion", "Cantidad", "Valor unitario", "IVA %", "IVA $", "Total linea"]
    if mode == "resumen":
        headers = ["Descripcion", "Cantidad", "Base", "IVA", "Total"]
    ws.append([])
    ws.append(headers)
    for cell in ws[8]:
        cell.font = Font(bold=True)

    for line in lines:
        if mode == "resumen":
            ws.append(
                [
                    line["descripcion"],
                    float(line["cantidad"] or 0),
                    float(line["base"] or 0),
                    float(line["iva"] or 0),
                    float(line["total"] or 0),
                ]
            )
        else:
            numero = "-"
            if line["numero"] is not None:
                numero = f"{line['prefijo']}-{int(line['numero']):06d}"
            ws.append(
                [
                    line["fecha"],
                    format_tipo_display(line["tipo"]),
                    numero,
                    line["descripcion"],
                    float(line["cantidad"] or 0),
                    float(line["precio_unitario"] or 0),
                    float((line["iva_rate"] or 0) * 100),
                    float(line["iva"] or 0),
                    float(line["total"] or 0),
                ]
            )

    total_row = ws.max_row + 1
    if mode == "resumen":
        ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
        ws.cell(row=total_row, column=2, value=float(totals["cantidad"])).font = Font(bold=True)
        ws.cell(row=total_row, column=3, value=float(totals["base"])).font = Font(bold=True)
        ws.cell(row=total_row, column=4, value=float(totals["iva"])).font = Font(bold=True)
        ws.cell(row=total_row, column=5, value=float(totals["total"])).font = Font(bold=True)
        ws.cell(row=total_row + 1, column=1, value="Cantidad documentos").font = Font(bold=True)
        ws.cell(row=total_row + 1, column=2, value=int(count_docs)).font = Font(bold=True)
    else:
        ws.cell(row=total_row, column=4, value="TOTAL").font = Font(bold=True)
        ws.cell(row=total_row, column=5, value=float(totals["cantidad"])).font = Font(bold=True)
        ws.cell(row=total_row, column=8, value=float(totals["iva"])).font = Font(bold=True)
        ws.cell(row=total_row, column=9, value=float(totals["total"])).font = Font(bold=True)
        ws.cell(row=total_row + 1, column=4, value="Cantidad documentos").font = Font(bold=True)
        ws.cell(row=total_row + 1, column=5, value=int(count_docs)).font = Font(bold=True)

    money_format = "#,##0.00"
    for row in ws.iter_rows(min_row=9, max_row=ws.max_row):
        for cell in row:
            if mode == "resumen":
                if cell.column in (3, 4, 5):
                    cell.number_format = money_format
                    cell.alignment = Alignment(horizontal="right")
                elif cell.column == 2:
                    cell.number_format = "0.00"
                    cell.alignment = Alignment(horizontal="right")
            else:
                if cell.column in (6, 8, 9):
                    cell.number_format = money_format
                    cell.alignment = Alignment(horizontal="right")
                elif cell.column == 7:
                    cell.number_format = "0.00"
                    cell.alignment = Alignment(horizontal="right")

    widths = [12, 10, 16, 36, 10, 16, 10, 16, 16]
    if mode == "resumen":
        widths = [36, 12, 16, 16, 16]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def build_informe_ventas_pdf(lines, header, totals, count_docs, mode, report):
    buffer = BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle("Title", parent=styles["Heading1"], alignment=1, fontSize=16, leading=18)
    normal = ParagraphStyle("Normal", parent=styles["Normal"], fontSize=10, leading=14)
    desc_style = ParagraphStyle("Desc", parent=styles["Normal"], fontSize=8, leading=10)

    story = []
    report_label = "Ventas por Cliente" if report == "ventas_cliente" else "Ventas por Producto"
    primary_label = "Cliente" if report == "ventas_cliente" else "Producto"
    secondary_label = "Producto" if report == "ventas_cliente" else "Cliente"

    story.append(Paragraph(report_label, title))
    story.append(Spacer(1, 6))
    story.append(Paragraph(f"<b>Emisor:</b> {header['emisor']}", normal))
    story.append(Paragraph(f"<b>Cliente:</b> {header['cliente']}", normal))
    story.append(Paragraph(f"<b>Rango de fechas:</b> {header['fechas']}", normal))
    story.append(Paragraph(f"<b>Tipo:</b> {header['tipo']}", normal))
    story.append(Spacer(1, 10))

    if mode == "resumen":
        table_data = [[primary_label, secondary_label, "Cantidad", "Valor total"]]
    else:
        table_data = [[primary_label, secondary_label, "Documento", "Cantidad", "Valor unitario", "Valor total"]]
    subtotal_rows = []
    for line in lines:
        if line.get("is_subtotal"):
            if mode == "resumen":
                table_data.append(
                    [
                        Paragraph(f"Subtotal {line['primary']}", desc_style),
                        Paragraph(line["secondary"] or "", desc_style),
                        f"{line['cantidad']:.2f}",
                        currency_cop(line["valor_total"] or 0),
                    ]
                )
            else:
                table_data.append(
                    [
                        Paragraph(f"Subtotal {line['primary']}", desc_style),
                        Paragraph(line["secondary"] or "", desc_style),
                        "",
                        f"{line['cantidad']:.2f}",
                        "",
                        currency_cop(line["valor_total"] or 0),
                    ]
                )
            subtotal_rows.append(len(table_data) - 1)
        else:
            if mode == "resumen":
                table_data.append(
                    [
                        Paragraph(str(line["primary"] or "-"), desc_style),
                        Paragraph(str(line["secondary"] or ""), desc_style),
                        f"{line['cantidad']:.2f}",
                        currency_cop(line["valor_total"] or 0),
                    ]
                )
            else:
                table_data.append(
                    [
                        Paragraph(str(line["primary"] or "-"), desc_style),
                        Paragraph(str(line["secondary"] or ""), desc_style),
                        line["documento"],
                        f"{line['cantidad']:.2f}",
                        currency_cop(line["valor_unitario"] or 0),
                        currency_cop(line["valor_total"] or 0),
                    ]
                )
    table_data.append(
        [Paragraph("TOTAL", desc_style), "", f"{totals['cantidad']:.2f}", currency_cop(totals["total"] or 0)]
        if mode == "resumen"
        else [
            Paragraph("TOTAL", desc_style),
            "",
            "",
            f"{totals['cantidad']:.2f}",
            "",
            currency_cop(totals["total"] or 0),
        ]
    )
    total_row = len(table_data) - 1

    col_widths = (
        [4.5 * cm, 4.5 * cm, 2.0 * cm, 3.0 * cm]
        if mode == "resumen"
        else [3.6 * cm, 3.6 * cm, 2.2 * cm, 1.8 * cm, 2.6 * cm, 3.0 * cm]
    )
    table = Table(table_data, repeatRows=1, colWidths=col_widths)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0f4f8")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.lightgrey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (1, -1), "LEFT"),
        ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("FONTNAME", (0, total_row), (-1, total_row), "Helvetica-Bold"),
    ]
    for row_idx in subtotal_rows:
        style.append(("FONTNAME", (0, row_idx), (-1, row_idx), "Helvetica-Bold"))
    table.setStyle(TableStyle(style))

    story.append(table)
    story.append(Spacer(1, 12))
    story.append(Paragraph(f"<b>Cantidad de documentos:</b> {count_docs}", normal))

    pdf.build(story)
    buffer.seek(0)
    return buffer


def build_informe_ventas_excel(lines, header, totals, count_docs, mode, report):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"

    report_label = "Ventas por Cliente" if report == "ventas_cliente" else "Ventas por Producto"
    primary_label = "Cliente" if report == "ventas_cliente" else "Producto"
    secondary_label = "Producto" if report == "ventas_cliente" else "Cliente"

    ws["A1"] = report_label
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generado: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A3"] = f"Emisor: {header['emisor']}"
    ws["A4"] = f"Cliente: {header['cliente']}"
    ws["A5"] = f"Rango de fechas: {header['fechas']}"
    ws["A6"] = f"Tipo: {header['tipo']}"

    headers = (
        [primary_label, secondary_label, "Cantidad", "Valor total"]
        if mode == "resumen"
        else [primary_label, secondary_label, "Documento", "Cantidad", "Valor unitario", "Valor total"]
    )
    ws.append([])
    ws.append(headers)
    for cell in ws[8]:
        cell.font = Font(bold=True)

    for line in lines:
        if line.get("is_subtotal"):
            row = [
                f"Subtotal {line['primary']}",
                line["secondary"] or "",
                float(line["cantidad"] or 0),
                float(line["valor_total"] or 0),
            ]
            if mode != "resumen":
                row = [
                    f"Subtotal {line['primary']}",
                    line["secondary"] or "",
                    "",
                    float(line["cantidad"] or 0),
                    "",
                    float(line["valor_total"] or 0),
                ]
            ws.append(row)
            for cell in ws[ws.max_row]:
                cell.font = Font(bold=True)
        else:
            if mode == "resumen":
                ws.append(
                    [
                        line["primary"],
                        line["secondary"] or "",
                        float(line["cantidad"] or 0),
                        float(line["valor_total"] or 0),
                    ]
                )
            else:
                ws.append(
                    [
                        line["primary"],
                        line["secondary"] or "",
                        line["documento"],
                        float(line["cantidad"] or 0),
                        float(line["valor_unitario"] or 0),
                        float(line["valor_total"] or 0),
                    ]
                )

    total_row = ws.max_row + 1
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    if mode == "resumen":
        ws.cell(row=total_row, column=3, value=float(totals["cantidad"])).font = Font(bold=True)
        ws.cell(row=total_row, column=4, value=float(totals["total"])).font = Font(bold=True)
    else:
        ws.cell(row=total_row, column=4, value=float(totals["cantidad"])).font = Font(bold=True)
        ws.cell(row=total_row, column=6, value=float(totals["total"])).font = Font(bold=True)
    ws.cell(row=total_row + 1, column=1, value="Cantidad documentos").font = Font(bold=True)
    ws.cell(row=total_row + 1, column=2, value=int(count_docs)).font = Font(bold=True)

    money_format = "#,##0.00"
    for row in ws.iter_rows(min_row=9, max_row=ws.max_row):
        for cell in row:
            if mode == "resumen":
                if cell.column in (4,):
                    cell.number_format = money_format
                    cell.alignment = Alignment(horizontal="right")
                elif cell.column == 3:
                    cell.number_format = "0.00"
                    cell.alignment = Alignment(horizontal="right")
            else:
                if cell.column in (5, 6):
                    cell.number_format = money_format
                    cell.alignment = Alignment(horizontal="right")
                elif cell.column == 4:
                    cell.number_format = "0.00"
                    cell.alignment = Alignment(horizontal="right")

    widths = [24, 24, 12, 16] if mode == "resumen" else [20, 24, 16, 12, 16, 16]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
init_db()


@app.get("/informes/documentos", response_class=HTMLResponse)
def informes_documentos(request: Request):
    conn = get_db()
    cur = conn.cursor()
    u = request.session.get("user") or {}
    if u.get("role") == "cliente_receptor":
        cur.execute("SELECT * FROM emisores WHERE id = ? ORDER BY nombre", (u.get("locked_emisor_id"),))
        emisores = cur.fetchall()
        cur.execute("SELECT * FROM receptores WHERE id = ? ORDER BY nombre", (u.get("locked_receptor_id"),))
        receptores = cur.fetchall()
    else:
        u = request.session.get("user") or {}
        if u.get("role") == "cliente_receptor":
            cur.execute("SELECT * FROM emisores WHERE id = ? ORDER BY nombre", (u.get("locked_emisor_id"),))
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores WHERE id = ? ORDER BY nombre", (u.get("locked_receptor_id"),))
            receptores = cur.fetchall()
        else:
            cur.execute("SELECT * FROM emisores ORDER BY nombre")
            emisores = cur.fetchall()
            cur.execute("SELECT * FROM receptores ORDER BY nombre")
            receptores = cur.fetchall()

    where, params, filters = parse_document_filters(request.query_params, default_estado=None)
    report = (request.query_params.get("report") or "ventas_cliente").lower()
    if report not in ("ventas_cliente", "ventas_producto"):
        report = "ventas_cliente"
    mode = (request.query_params.get("mode") or "detalle").lower()
    if mode not in ("detalle", "resumen"):
        mode = "detalle"
    lines_raw = fetch_report_lines(conn, where, params)
    doc_ids = {line["documento_id"] for line in lines_raw}
    lines, totals_ventas = build_ventas_lines(lines_raw, mode, report)
    total_qty = totals_ventas["cantidad"]
    total_base = totals_ventas["total"]
    total_iva = 0.0
    total_consulta = totals_ventas["total"]

    emisor_text = "Todos"
    cliente_text = "Todos"
    if filters.get("emisor_id"):
        cur.execute("SELECT nombre, identificacion, tipo_identificacion FROM emisores WHERE id = ?", (filters["emisor_id"],))
        row = cur.fetchone()
        if row:
            emisor_text = f"{row['nombre']} ({format_id_with_type(row['tipo_identificacion'], row['identificacion'])})"
    if filters.get("receptor_id"):
        cur.execute("SELECT nombre, identificacion, tipo_identificacion FROM receptores WHERE id = ?", (filters["receptor_id"],))
        row = cur.fetchone()
        if row:
            cliente_text = f"{row['nombre']} ({format_id_with_type(row['tipo_identificacion'], row['identificacion'])})"

    fecha_text = "Todas"
    if filters.get("fecha_desde") or filters.get("fecha_hasta"):
        desde = filters.get("fecha_desde") or "-"
        hasta = filters.get("fecha_hasta") or "-"
        fecha_text = f"{desde} a {hasta}"

    tipo_text = filters.get("tipo_documento") or "Todos"
    conn.close()

    filters["mode"] = mode
    filters["report"] = report
    query_string = build_filters_query_string(filters)
    return templates.TemplateResponse(
        "informes_documentos.html",
        {
            "request": request,
            "emisores": emisores,
            "receptores": receptores,
            "lines": lines,
            "filters": filters,
            "totals": {"base": total_base, "iva": total_iva, "total": total_consulta, "cantidad": total_qty},
            "total_count": len(doc_ids),
            "header": {
                "emisor": emisor_text,
                "cliente": cliente_text,
                "fechas": fecha_text,
                "tipo": tipo_text,
            },
            "mode": mode,
            "report": report,
            "query_string": query_string,
        },
    )


@app.get("/informes/documentos/pdf")
def informes_documentos_pdf(request: Request, inline: bool = False):
    conn = get_db()
    where, params, filters = parse_document_filters(request.query_params, default_estado=None)
    mode = (request.query_params.get("mode") or "detalle").lower()
    if mode not in ("detalle", "resumen"):
        mode = "detalle"
    lines_raw = fetch_report_lines(conn, where, params)
    doc_ids = {line["documento_id"] for line in lines_raw}
    if mode == "resumen":
        grouped = {}
        for line in lines_raw:
            key = (line["descripcion"] or "")
            item = grouped.get(key) or {
                "descripcion": key,
                "cantidad": 0.0,
                "base": 0.0,
                "iva": 0.0,
                "total": 0.0,
            }
            item["cantidad"] += float(line["cantidad"] or 0)
            item["base"] += float(line["subtotal"] or 0)
            item["iva"] += float(line["iva"] or 0)
            item["total"] += float(line["total"] or 0)
            grouped[key] = item
        lines = list(grouped.values())
        total_base = sum(item["base"] for item in lines)
        total_iva = sum(item["iva"] for item in lines)
        total_consulta = sum(item["total"] for item in lines)
        total_qty = sum(item["cantidad"] for item in lines)
    else:
        lines = lines_raw
        total_base = sum(float(line["subtotal"] or 0) for line in lines)
        total_iva = sum(float(line["iva"] or 0) for line in lines)
        total_consulta = sum(float(line["total"] or 0) for line in lines)
        total_qty = sum(float(line["cantidad"] or 0) for line in lines)

    emisor_text = "Todos"
    cliente_text = "Todos"
    if filters.get("emisor_id"):
        cur = conn.cursor()
        cur.execute("SELECT nombre, identificacion, tipo_identificacion FROM emisores WHERE id = ?", (filters["emisor_id"],))
        row = cur.fetchone()
        if row:
            emisor_text = f"{row['nombre']} ({format_id_with_type(row['tipo_identificacion'], row['identificacion'])})"
    if filters.get("receptor_id"):
        cur = conn.cursor()
        cur.execute("SELECT nombre, identificacion, tipo_identificacion FROM receptores WHERE id = ?", (filters["receptor_id"],))
        row = cur.fetchone()
        if row:
            cliente_text = f"{row['nombre']} ({format_id_with_type(row['tipo_identificacion'], row['identificacion'])})"

    fecha_text = "Todas"
    if filters.get("fecha_desde") or filters.get("fecha_hasta"):
        desde = filters.get("fecha_desde") or "-"
        hasta = filters.get("fecha_hasta") or "-"
        fecha_text = f"{desde} a {hasta}"

    tipo_text = filters.get("tipo_documento") or "Todos"
    conn.close()

    pdf_buffer = build_informe_documentos_pdf(
        lines,
        {"emisor": emisor_text, "cliente": cliente_text, "fechas": fecha_text, "tipo": tipo_text},
        {"base": total_base, "iva": total_iva, "total": total_consulta, "cantidad": total_qty},
        len(doc_ids),
        mode,
    )
    filename = "Informe_documentos.pdf"
    disposition = "inline" if inline else "attachment"
    headers = {"Content-Disposition": f'{disposition}; filename="{filename}"'}
    return StreamingResponse(pdf_buffer, media_type="application/pdf", headers=headers)


@app.get("/informes/documentos/pdf/preview")
def informes_documentos_pdf_preview(request: Request):
    return informes_documentos_pdf(request, inline=True)


@app.get("/informes/documentos/excel")
def informes_documentos_excel(request: Request):
    conn = get_db()
    where, params, filters = parse_document_filters(request.query_params, default_estado=None)
    mode = (request.query_params.get("mode") or "detalle").lower()
    if mode not in ("detalle", "resumen"):
        mode = "detalle"
    lines_raw = fetch_report_lines(conn, where, params)
    doc_ids = {line["documento_id"] for line in lines_raw}
    if mode == "resumen":
        grouped = {}
        for line in lines_raw:
            key = (line["descripcion"] or "")
            item = grouped.get(key) or {
                "descripcion": key,
                "cantidad": 0.0,
                "base": 0.0,
                "iva": 0.0,
                "total": 0.0,
            }
            item["cantidad"] += float(line["cantidad"] or 0)
            item["base"] += float(line["subtotal"] or 0)
            item["iva"] += float(line["iva"] or 0)
            item["total"] += float(line["total"] or 0)
            grouped[key] = item
        lines = list(grouped.values())
        total_base = sum(item["base"] for item in lines)
        total_iva = sum(item["iva"] for item in lines)
        total_consulta = sum(item["total"] for item in lines)
        total_qty = sum(item["cantidad"] for item in lines)
    else:
        lines = lines_raw
        total_base = sum(float(line["subtotal"] or 0) for line in lines)
        total_iva = sum(float(line["iva"] or 0) for line in lines)
        total_consulta = sum(float(line["total"] or 0) for line in lines)
        total_qty = sum(float(line["cantidad"] or 0) for line in lines)

    emisor_text = "Todos"
    cliente_text = "Todos"
    if filters.get("emisor_id"):
        cur = conn.cursor()
        cur.execute("SELECT nombre, identificacion, tipo_identificacion FROM emisores WHERE id = ?", (filters["emisor_id"],))
        row = cur.fetchone()
        if row:
            emisor_text = f"{row['nombre']} ({format_id_with_type(row['tipo_identificacion'], row['identificacion'])})"
    if filters.get("receptor_id"):
        cur = conn.cursor()
        cur.execute("SELECT nombre, identificacion, tipo_identificacion FROM receptores WHERE id = ?", (filters["receptor_id"],))
        row = cur.fetchone()
        if row:
            cliente_text = f"{row['nombre']} ({format_id_with_type(row['tipo_identificacion'], row['identificacion'])})"

    fecha_text = "Todas"
    if filters.get("fecha_desde") or filters.get("fecha_hasta"):
        desde = filters.get("fecha_desde") or "-"
        hasta = filters.get("fecha_hasta") or "-"
        fecha_text = f"{desde} a {hasta}"

    tipo_text = filters.get("tipo_documento") or "Todos"
    conn.close()

    output = build_informe_documentos_excel(
        lines,
        {"emisor": emisor_text, "cliente": cliente_text, "fechas": fecha_text, "tipo": tipo_text},
        {"base": total_base, "iva": total_iva, "total": total_consulta, "cantidad": total_qty},
        len(doc_ids),
        mode,
    )
    filename = "Informe_documentos.xlsx"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@app.get("/informes/ventas/pdf")
def informes_ventas_pdf(request: Request, inline: bool = False):
    conn = get_db()
    where, params, filters = parse_document_filters(request.query_params, default_estado=None)
    lines_raw = fetch_report_lines(conn, where, params)
    doc_ids = {line["documento_id"] for line in lines_raw}
    mode = (request.query_params.get("mode") or "detalle").lower()
    if mode not in ("detalle", "resumen"):
        mode = "detalle"
    report = (request.query_params.get("report") or "ventas_cliente").lower()
    if report not in ("ventas_cliente", "ventas_producto"):
        report = "ventas_cliente"
    lines, totals_ventas = build_ventas_lines(lines_raw, mode, report)

    emisor_text = "Todos"
    cliente_text = "Todos"
    if filters.get("emisor_id"):
        cur = conn.cursor()
        cur.execute("SELECT nombre, identificacion, tipo_identificacion FROM emisores WHERE id = ?", (filters["emisor_id"],))
        row = cur.fetchone()
        if row:
            emisor_text = f"{row['nombre']} ({format_id_with_type(row['tipo_identificacion'], row['identificacion'])})"
    if filters.get("receptor_id"):
        cur = conn.cursor()
        cur.execute("SELECT nombre, identificacion, tipo_identificacion FROM receptores WHERE id = ?", (filters["receptor_id"],))
        row = cur.fetchone()
        if row:
            cliente_text = f"{row['nombre']} ({format_id_with_type(row['tipo_identificacion'], row['identificacion'])})"

    fecha_text = "Todas"
    if filters.get("fecha_desde") or filters.get("fecha_hasta"):
        desde = filters.get("fecha_desde") or "-"
        hasta = filters.get("fecha_hasta") or "-"
        fecha_text = f"{desde} a {hasta}"

    tipo_text = filters.get("tipo_documento") or "Todos"
    conn.close()

    pdf_buffer = build_informe_ventas_pdf(
        lines,
        {"emisor": emisor_text, "cliente": cliente_text, "fechas": fecha_text, "tipo": tipo_text},
        totals_ventas,
        len(doc_ids),
        mode,
        report,
    )
    filename = "Ventas_por_Cliente.pdf" if report == "ventas_cliente" else "Ventas_por_Producto.pdf"
    disposition = "inline" if inline else "attachment"
    headers = {"Content-Disposition": f'{disposition}; filename="{filename}"'}
    return StreamingResponse(pdf_buffer, media_type="application/pdf", headers=headers)


@app.get("/informes/ventas/pdf/preview")
def informes_ventas_pdf_preview(request: Request):
    return informes_ventas_pdf(request, inline=True)


@app.get("/informes/ventas/excel")
def informes_ventas_excel(request: Request):
    conn = get_db()
    where, params, filters = parse_document_filters(request.query_params, default_estado=None)
    lines_raw = fetch_report_lines(conn, where, params)
    doc_ids = {line["documento_id"] for line in lines_raw}
    mode = (request.query_params.get("mode") or "detalle").lower()
    if mode not in ("detalle", "resumen"):
        mode = "detalle"
    report = (request.query_params.get("report") or "ventas_cliente").lower()
    if report not in ("ventas_cliente", "ventas_producto"):
        report = "ventas_cliente"
    lines, totals_ventas = build_ventas_lines(lines_raw, mode, report)

    emisor_text = "Todos"
    cliente_text = "Todos"
    if filters.get("emisor_id"):
        cur = conn.cursor()
        cur.execute("SELECT nombre, identificacion, tipo_identificacion FROM emisores WHERE id = ?", (filters["emisor_id"],))
        row = cur.fetchone()
        if row:
            emisor_text = f"{row['nombre']} ({format_id_with_type(row['tipo_identificacion'], row['identificacion'])})"
    if filters.get("receptor_id"):
        cur = conn.cursor()
        cur.execute("SELECT nombre, identificacion, tipo_identificacion FROM receptores WHERE id = ?", (filters["receptor_id"],))
        row = cur.fetchone()
        if row:
            cliente_text = f"{row['nombre']} ({format_id_with_type(row['tipo_identificacion'], row['identificacion'])})"

    fecha_text = "Todas"
    if filters.get("fecha_desde") or filters.get("fecha_hasta"):
        desde = filters.get("fecha_desde") or "-"
        hasta = filters.get("fecha_hasta") or "-"
        fecha_text = f"{desde} a {hasta}"

    tipo_text = filters.get("tipo_documento") or "Todos"
    conn.close()

    output = build_informe_ventas_excel(
        lines,
        {"emisor": emisor_text, "cliente": cliente_text, "fechas": fecha_text, "tipo": tipo_text},
        totals_ventas,
        len(doc_ids),
        mode,
        report,
    )
    filename = "Ventas_por_Cliente.xlsx" if report == "ventas_cliente" else "Ventas_por_Producto.xlsx"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )
@app.get("/login", response_class=HTMLResponse)
def login_get(request: Request, next: str = "/"):
    return templates.TemplateResponse(
        "login.html",
        {"request": request, "error": None, "next": next}
    )

@app.post("/login", response_class=HTMLResponse)
def login_post(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    next: str = Form("/"),
):
    user = get_user_by_username(username.strip())

    if (not user) or (user["is_active"] != 1) or (not pwd_context.verify(password, user["password_hash"])):
        return templates.TemplateResponse(
            "login.html",
            {"request": request, "error": "Usuario o contraseña incorrectos", "next": next},
            status_code=401,
        )
    request.session["user"] = {
        "id": user["id"],
        "username": user["username"],
        "role": user["role"],
        "locked_emisor_id": user["locked_emisor_id"],
        "locked_receptor_id": user["locked_receptor_id"],
    }
    return RedirectResponse(url=next or "/", status_code=302)

@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse(url="/login", status_code=302)

